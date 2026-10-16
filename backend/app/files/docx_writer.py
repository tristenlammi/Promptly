"""A document's HTML → .docx, on the server (WP-10.4, Part C P8).

The editor's own .docx export runs in the browser (``lib/office/docxExport.ts``,
WP-4.3). Exporting a whole library can't wait on a browser, so this writes
the same document with python-docx from the HTML the PDF export already
uses (``document_render``'s output, images inlined as ``data:`` URIs).
"Good enough to open in Word and LibreOffice", like the browser's: headings,
paragraphs and alignment, bold/italic/underline/strike/code/sub/super,
colour and highlight, links, bullet / numbered / task lists (nested),
quotes, code blocks, tables, images, rules, page breaks and the page size
and margins. Anything it doesn't know becomes its text.
"""
from __future__ import annotations

import base64
import binascii
import io
import re
from dataclasses import dataclass, replace
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_BLOCKS = {
    "p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "blockquote", "pre", "hr", "table",
    "div", "section", "article", "details", "summary", "figure", "figcaption", "header", "footer", "main",
}
_MAX_IMAGE_WIDTH = Inches(6)
_CODE_FONT = "Consolas"


@dataclass(frozen=True)
class _Marks:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    code: bool = False
    sub: bool = False
    sup: bool = False
    color: str | None = None
    size: float | None = None
    highlight: bool = False
    href: str | None = None
    # A suggestion (suggestion mode) → a Word tracked change: (ins|del, author, ISO time).
    suggest: tuple[str, str, str] | None = None


def _style(el: Any) -> dict[str, str]:
    out = {}
    for part in (el.get("style") or "").split(";"):
        k, _, v = part.partition(":")
        if k.strip() and v.strip():
            out[k.strip().lower()] = v.strip()
    return out


def _hex(value: str | None) -> str | None:
    if not value:
        return None
    m = re.fullmatch(r"#([0-9a-f]{6}|[0-9a-f]{3})", value.strip().lower())
    if m:
        h = m.group(1)
        return h if len(h) == 6 else "".join(c * 2 for c in h)
    m = re.fullmatch(r"rgba?\((\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,\s*[\d.]+)?\)", value.strip().lower())
    if m:
        return "".join(f"{min(int(x), 255):02x}" for x in m.groups())
    return None


class _Writer:
    def __init__(self, doc: Any) -> None:
        self.doc = doc

    # -------------------------------------------------------------- blocks

    def blocks(self, parent: Any, container: Any, *, style: str | None = None, depth: int = 0) -> None:
        """Write ``parent``'s children into ``container`` (the document or a
        table cell). Loose inline content between blocks gathers into one
        paragraph."""
        pending: list[Any] = []
        lead = (parent.text or "") if parent.tag not in ("ul", "ol") else ""

        def flush() -> None:
            nonlocal lead
            if pending or lead.strip():
                p = container.add_paragraph(style=style)
                if lead:
                    self._text(p, lead, _Marks())
                for node in pending:
                    self.inline(p, node, _Marks())
                pending.clear()
            lead = ""

        for child in parent:
            tag = child.tag if isinstance(child.tag, str) else ""
            if tag in _BLOCKS or (tag == "img" and parent.tag not in ("p",)):
                flush()
                self.block(child, container, style=style, depth=depth)
                if child.tail and child.tail.strip():
                    lead = child.tail
            else:
                pending.append(child)
        flush()

    def block(self, el: Any, container: Any, *, style: str | None, depth: int) -> None:
        tag = el.tag
        kind = el.get("data-type") or ""
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6") and container is self.doc:
            p = self.doc.add_heading(level=int(tag[1]))
            self._para(p, el)
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            p = container.add_paragraph(style=style)
            self._para(p, el, _Marks(bold=True))
        elif tag in ("p", "summary"):
            p = container.add_paragraph(style=style)
            self._para(p, el, _Marks(bold=tag == "summary"))
        elif tag in ("ul", "ol"):
            self._list(el, container, ordered=tag == "ol", task=kind == "taskList", depth=depth)
        elif tag == "blockquote":
            self.blocks(el, container, style="Quote" if container is self.doc else style, depth=depth)
        elif tag == "pre":
            p = container.add_paragraph(style=style)
            lines = el.text_content().split("\n")
            for i, line in enumerate(lines):
                run = p.add_run(line)
                run.font.name = _CODE_FONT
                run.font.size = Pt(9.5)
                if i < len(lines) - 1:
                    run.add_break()
        elif tag == "hr":
            self._rule(container.add_paragraph(style=style))
        elif tag == "table":
            self._table(el, container)
        elif tag == "img":
            p = container.add_paragraph(style=style)
            self._image(p, el)
        elif kind == "page-break" and container is self.doc:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "block-math":
            p = container.add_paragraph(style=style)
            run = p.add_run(el.get("data-latex") or el.text_content())
            run.font.name = _CODE_FONT
        else:
            # div, details, sections, a callout, columns: their contents.
            self.blocks(el, container, style=style, depth=depth)

    def _para(self, p: Any, el: Any, marks: _Marks | None = None) -> None:
        align = _style(el).get("text-align")
        if align in _ALIGN:
            p.alignment = _ALIGN[align]
        marks = marks or _Marks()
        if el.text:
            self._text(p, el.text, marks)
        for child in el:
            self.inline(p, child, marks)

    def _list(self, el: Any, container: Any, *, ordered: bool, task: bool, depth: int) -> None:
        base = "List Number" if ordered else "List Bullet"
        style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
        if container is not self.doc or task:
            style = None  # a task's box is its marker
        n = 0
        for li in el:
            if not isinstance(li.tag, str) or li.tag != "li":
                continue
            n += 1
            p = container.add_paragraph(style=style)
            if style is None and not task:
                # In a table cell: no list styles there, so the marker is text.
                p.paragraph_format.left_indent = Inches(0.25 * depth)
                p.add_run(f"{n}. " if ordered else "• ")
            if task:
                p.paragraph_format.left_indent = Inches(0.25 * depth)
                p.add_run("☑ " if li.get("data-checked") == "true" else "☐ ")
            # The item's first paragraph is its line; nested lists indent.
            first = True
            if li.text and li.text.strip():
                self._text(p, li.text, _Marks())
                first = False
            for child in li:
                ctag = child.tag if isinstance(child.tag, str) else ""
                if ctag in ("ul", "ol"):
                    self._list(child, container, ordered=ctag == "ol", task=child.get("data-type") == "taskList", depth=depth + 1)
                elif ctag in ("p", "div") and first:
                    self._para(p, child)
                    first = False
                elif ctag in _BLOCKS:
                    extra = container.add_paragraph()
                    extra.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                    self._para(extra, child)
                else:
                    self.inline(p, child, _Marks())
                    first = False

    def _table(self, el: Any, container: Any) -> None:
        rows = [tr for tr in el.iter("tr")]
        if not rows:
            return
        width = max(sum(int(c.get("colspan") or 1) for c in tr if c.tag in ("td", "th")) for tr in rows)
        if width == 0:
            return
        table = container.add_table(rows=len(rows), cols=width)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):
            pass
        for r, tr in enumerate(rows):
            c = 0
            for cell_el in tr:
                if cell_el.tag not in ("td", "th") or c >= width:
                    continue
                span = max(1, min(int(cell_el.get("colspan") or 1), width - c))
                cell = table.cell(r, c)
                if span > 1:
                    cell = cell.merge(table.cell(r, c + span - 1))
                first = cell.paragraphs[0]
                _CellAdapter(cell, first, bold=cell_el.tag == "th").fill(self, cell_el)
                c += span

    def _rule(self, p: Any) -> None:
        ppr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BFBFBF")):
            bottom.set(qn(k), v)
        border.append(bottom)
        ppr.append(border)

    # -------------------------------------------------------------- inline

    def inline(self, p: Any, el: Any, marks: _Marks) -> None:
        tag = el.tag if isinstance(el.tag, str) else ""
        m = marks
        if tag in ("strong", "b"):
            m = replace(m, bold=True)
        elif tag in ("em", "i"):
            m = replace(m, italic=True)
        elif tag == "u":
            m = replace(m, underline=True)
        elif tag in ("s", "del", "strike"):
            m = replace(m, strike=True)
        elif tag == "code":
            m = replace(m, code=True)
        elif tag == "sub":
            m = replace(m, sub=True)
        elif tag == "sup":
            m = replace(m, sup=True)
        elif tag == "mark":
            m = replace(m, highlight=True)
        elif tag == "a" and (el.get("href") or "").startswith(("http://", "https://", "mailto:")):
            m = replace(m, href=el.get("href"))
        if tag == "span" and el.get("data-suggest") in ("ins", "del"):
            m = replace(m, suggest=(el.get("data-suggest"), el.get("data-author") or "Promptly", el.get("data-at") or ""))
        if tag in ("span", "mark"):
            st = _style(el)
            colour = _hex(st.get("color"))
            if colour:
                m = replace(m, color=colour)
            size = re.fullmatch(r"([\d.]+)(px|pt)", st.get("font-size", ""))
            if size:
                pts = float(size.group(1)) * (0.75 if size.group(2) == "px" else 1)
                m = replace(m, size=max(4.0, min(pts, 96.0)))
        if tag == "br":
            p.add_run().add_break()
        elif tag == "img":
            self._image(p, el)
        elif el.get("data-type") == "inline-math":
            self._text(p, el.get("data-latex") or el.text_content(), replace(m, code=True))
        else:
            if el.text:
                self._text(p, el.text, m)
            for child in el:
                self.inline(p, child, m)
        if el.tail:
            self._text(p, el.tail, marks)

    def _text(self, p: Any, text: str, m: _Marks) -> None:
        text = re.sub(r"[\r\n\t ]+", " ", text)
        if not text:
            return
        run = p.add_run(text)
        f = run.font
        f.bold = m.bold or None
        f.italic = m.italic or None
        f.underline = m.underline or bool(m.href) or None
        f.strike = m.strike or None
        f.subscript = m.sub or None
        f.superscript = m.sup or None
        if m.code:
            f.name = _CODE_FONT
        if m.size:
            f.size = Pt(m.size)
        if m.color:
            f.color.rgb = RGBColor.from_string(m.color.upper())
        elif m.href:
            f.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        if m.highlight:
            f.highlight_color = WD_COLOR_INDEX.YELLOW
        node = run._r
        if m.href:
            r_id = p.part.relate_to(m.href, RT.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), r_id)
            link.append(run._r)  # moved out of the paragraph into the link
            p._p.append(link)
            node = link
        if m.suggest:
            self._tracked(p, node, run, m.suggest)

    def _tracked(self, p: Any, node: Any, run: Any, suggest: tuple[str, str, str]) -> None:
        """A suggestion as Word's own tracked change (``w:ins`` / ``w:del``),
        so it opens in Word as a change to accept or reject."""
        kind, author, at = suggest
        self._change_id = getattr(self, "_change_id", 0) + 1
        wrap = OxmlElement("w:ins" if kind == "ins" else "w:del")
        wrap.set(qn("w:id"), str(self._change_id))
        wrap.set(qn("w:author"), author[:120])
        stamp = re.match(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}", at or "")
        if stamp:
            wrap.set(qn("w:date"), stamp.group(0) + "Z")
        if kind == "del":
            # Deleted text is w:delText in Word.
            for t in run._r.findall(qn("w:t")):
                t.tag = qn("w:delText")
        node.addprevious(wrap)
        wrap.append(node)

    def _image(self, p: Any, el: Any) -> None:
        data = _data_uri(el.get("src") or "")
        if data is None:
            alt = el.get("alt") or ""
            if alt:
                self._text(p, f"[{alt}]", _Marks(italic=True))
            return
        try:
            from PIL import Image

            with Image.open(io.BytesIO(data)) as im:
                px_w = im.size[0]
                fmt = (im.format or "").upper()
                if fmt not in ("PNG", "JPEG", "GIF", "BMP"):
                    # Word reads only these: anything else goes in as PNG.
                    buf = io.BytesIO()
                    im.convert("RGBA").save(buf, "PNG")
                    data = buf.getvalue()
        except Exception:  # noqa: BLE001 — a picture that won't open is left out
            return
        want = el.get("width")
        px = int(want) if want and want.isdigit() else px_w
        width = min(Emu(int(px * 9525)), _MAX_IMAGE_WIDTH)  # 96 dpi
        try:
            p.add_run().add_picture(io.BytesIO(data), width=width if px_w else None)
        except Exception:  # noqa: BLE001
            return


class _CellAdapter:
    """A table cell: the first block goes into its existing paragraph."""

    def __init__(self, cell: Any, first: Any, *, bold: bool) -> None:
        self.cell, self.first, self.bold, self.used = cell, first, bold, False

    def add_paragraph(self, style: str | None = None) -> Any:  # noqa: ARG002 — cells keep the table's style
        if not self.used:
            self.used = True
            return self.first
        return self.cell.add_paragraph()

    def add_table(self, rows: int, cols: int) -> Any:
        return self.cell.add_table(rows=rows, cols=cols)

    def fill(self, writer: _Writer, el: Any) -> None:
        writer.blocks(el, self)
        if self.bold:
            for p in self.cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True


def _data_uri(src: str) -> bytes | None:
    m = re.match(r"data:image/[\w.+-]+;base64,(.*)", src, re.S)
    if not m:
        return None
    try:
        return base64.b64decode(m.group(1), validate=False)
    except (binascii.Error, ValueError):
        return None


def html_to_docx(html_text: str, *, title: str = "", page: Any = None) -> bytes:
    """The document as a .docx. ``page`` is its :class:`PageSetup`."""
    from lxml import html as lhtml

    doc = Document()
    if title:
        doc.core_properties.title = title[:255]
    if page is not None:
        section = doc.sections[0]
        w, h = page.paper_mm()
        section.page_width, section.page_height = Mm(w), Mm(h)
        if page.orientation == "landscape":
            from docx.enum.section import WD_ORIENT

            section.orientation = WD_ORIENT.LANDSCAPE
        section.top_margin, section.right_margin = Mm(page.margin_top), Mm(page.margin_right)
        section.bottom_margin, section.left_margin = Mm(page.margin_bottom), Mm(page.margin_left)
        if page.header:
            section.header.paragraphs[0].text = page.header
        if page.footer:
            section.footer.paragraphs[0].text = page.footer
    if (html_text or "").strip():
        root = lhtml.fragment_fromstring(html_text, create_parent="div")
        _Writer(doc).blocks(root, doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


__all__ = ["html_to_docx"]
