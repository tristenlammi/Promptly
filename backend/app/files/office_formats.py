"""Office files as uploads (WP-4.3, Part C §3.7).

Word, Excel, PowerPoint and OpenDocument files can be uploaded, previewed
and searched. Two guards and one job live here:

* :func:`check_zip_container` — every OOXML / ODF file is a zip, and a
  zip can be a bomb (a 207 KB document inflated to 885 MB in the format
  spike). Before anything parses one: at most 10,000 entries, 500 MB
  uncompressed, and a compression ratio of 100. Also confirms the zip
  really is the format its extension claims.
* :func:`extract_office_text` — the plain text for Drive search, RAG and
  "Ask about this", via python-docx / openpyxl / python-pptx (and the ODF
  ``content.xml`` through defusedxml). It runs on the Arq worker after an
  upload, never inside the request.

Converting to an editable Promptly document happens in the browser, from
the editor's own schema (``frontend/src/lib/office/``).
"""
from __future__ import annotations

import logging
import re
import zipfile
from pathlib import Path

logger = logging.getLogger("promptly.files.office_formats")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
ODT_MIME = "application/vnd.oasis.opendocument.text"
ODS_MIME = "application/vnd.oasis.opendocument.spreadsheet"
ODP_MIME = "application/vnd.oasis.opendocument.presentation"
RTF_MIME = "application/rtf"

# ext → (canonical MIME, aliases). Generic zip is an alias because some
# writers' files sniff as a plain zip; check_zip_container then confirms
# what's inside.
OFFICE_ALLOWLIST: dict[str, tuple[str, frozenset[str]]] = {
    ".docx": (DOCX_MIME, frozenset({"application/zip", "application/x-zip-compressed"})),
    ".xlsx": (XLSX_MIME, frozenset({"application/zip", "application/x-zip-compressed"})),
    ".pptx": (PPTX_MIME, frozenset({"application/zip", "application/x-zip-compressed"})),
    ".odt": (ODT_MIME, frozenset({"application/zip"})),
    ".ods": (ODS_MIME, frozenset({"application/zip"})),
    ".odp": (ODP_MIME, frozenset({"application/zip"})),
    ".rtf": (RTF_MIME, frozenset({"text/rtf", "text/plain"})),
}
# Legacy Office (OLE compound files): stored and downloadable; previewed
# only by the optional converter (Stage 12) — never parsed here. The
# three are aliases of each other: sniffers disagree on which is which.
_OLE = frozenset({"application/msword", "application/vnd.ms-excel", "application/vnd.ms-powerpoint",
                  "application/x-ole-storage", "application/cdfv2"})
LEGACY_OFFICE_ALLOWLIST: dict[str, tuple[str, frozenset[str]]] = {
    ".doc": ("application/msword", _OLE),
    ".xls": ("application/vnd.ms-excel", _OLE),
    ".ppt": ("application/vnd.ms-powerpoint", _OLE),
}
ZIP_OFFICE_EXTS = frozenset(e for e in OFFICE_ALLOWLIST if e != ".rtf")
OFFICE_EXTS = frozenset(OFFICE_ALLOWLIST)

MAX_ZIP_ENTRIES = 10_000
MAX_ZIP_UNCOMPRESSED = 500 * 1024 * 1024
MAX_ZIP_RATIO = 100
# Where each format's main part lives — a zip without it isn't that format.
_REQUIRED_PART = {
    ".docx": "word/document.xml",
    ".xlsx": "xl/workbook.xml",
    ".pptx": "ppt/presentation.xml",
    ".odt": "content.xml",
    ".ods": "content.xml",
    ".odp": "content.xml",
}

TEXT_LIMIT = 256 * 1024


def office_ext(filename: str | None) -> str | None:
    ext = Path(filename or "").suffix.lower()
    return ext if ext in OFFICE_EXTS else None


def check_zip_container(path: Path, ext: str) -> None:
    """Refuse a zip that would blow up when opened, or isn't what it says.
    Raises :class:`app.files.safety.UnsafeUploadError`."""
    from app.files.safety import UnsafeUploadError

    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise UnsafeUploadError("zip_bomb", "This file has too many parts to open safely.")
            total = sum(i.file_size for i in infos)
            packed = sum(i.compress_size for i in infos) or 1
            if total > MAX_ZIP_UNCOMPRESSED or total / packed > MAX_ZIP_RATIO:
                raise UnsafeUploadError(
                    "zip_bomb", "This file expands to far more than its size — it can't be opened safely."
                )
            names = {i.filename for i in infos}
    except zipfile.BadZipFile:
        raise UnsafeUploadError("mime_mismatch", "This file is damaged or isn't really an Office document.")
    required = _REQUIRED_PART.get(ext)
    if required and required not in names:
        raise UnsafeUploadError(
            "mime_mismatch", f"This doesn't look like a {ext[1:].upper()} file (it has no {required})."
        )


def _cap(parts: list[str]) -> str:
    out: list[str] = []
    size = 0
    for p in parts:
        p = p.strip()
        if not p:
            continue
        size += len(p) + 1
        if size > TEXT_LIMIT:
            out.append("[truncated]")
            break
        out.append(p)
    return "\n".join(out)


def _docx_text(path: Path) -> str:
    import docx  # python-docx; its parser doesn't resolve entities

    d = docx.Document(str(path))
    parts: list[str] = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append("".join(t.text or "" for t in block.iter() if t.tag.endswith("}t")))
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iterchildren():
                        if cell.tag.endswith("}tc"):
                            cells.append(" ".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t")).strip())
                    parts.append(" | ".join(cells))
    return _cap(parts)


def _xlsx_text(path: Path) -> str:
    import openpyxl  # uses defusedxml when it's installed (it is)

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"## {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 10_000:
                    break
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals:
                    parts.append(" | ".join(vals))
    finally:
        wb.close()
    return _cap(parts)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        parts.append(f"## Slide {n}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                parts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
    return _cap(parts)


def _odf_text(path: Path) -> str:
    from defusedxml import ElementTree as DET

    with zipfile.ZipFile(path) as zf:
        raw = zf.read("content.xml")
    root = DET.fromstring(raw, forbid_dtd=True)
    parts: list[str] = []
    for el in root.iter():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag in ("p", "h"):
            parts.append("".join(el.itertext()))
    return _cap(parts)


_RTF_CONTROL = re.compile(r"\\[a-z]+-?\d* ?|\\'[0-9a-f]{2}|[{}]", re.I)


def _rtf_text(path: Path) -> str:
    raw = path.read_bytes()[: TEXT_LIMIT * 4].decode("latin-1", errors="replace")
    # Drop the font / colour / style tables, then control words.
    raw = re.sub(r"\{\\\*[^{}]*\}|\{\\(fonttbl|colortbl|stylesheet|info)[^{}]*(\{[^{}]*\}[^{}]*)*\}", "", raw)
    return _cap(_RTF_CONTROL.sub("", raw).replace("\\par", "\n").splitlines())


def extract_office_text(path: Path, ext: str) -> str | None:
    """The file's plain text, or None. Never raises — a damaged file is just
    not searchable. Zip formats are re-checked before parsing."""
    try:
        if ext in ZIP_OFFICE_EXTS:
            check_zip_container(path, ext)
        if ext == ".docx":
            return _docx_text(path) or None
        if ext == ".xlsx":
            return _xlsx_text(path) or None
        if ext == ".pptx":
            return _pptx_text(path) or None
        if ext in (".odt", ".ods", ".odp"):
            return _odf_text(path) or None
        if ext == ".rtf":
            return _rtf_text(path) or None
    except Exception:  # noqa: BLE001 — extraction is best-effort
        logger.warning("office text extraction failed for %s", path.name, exc_info=True)
    return None


async def extract_and_store(file_id) -> bool:
    """Arq job body: read an uploaded Office file's text into
    ``content_text`` (Drive search + RAG + chat). True when text was found."""
    import asyncio
    import uuid

    # The worker process doesn't load the whole app: register the tables
    # ``files`` points at (users) before touching it.
    import app.auth.models  # noqa: F401
    from app.database import SessionLocal
    from app.files.models import UserFile
    from app.files.storage import absolute_path

    async with SessionLocal() as db:
        row = await db.get(UserFile, uuid.UUID(str(file_id)))
        if row is None or row.trashed_at is not None:
            return False
        ext = office_ext(row.filename)
        if ext is None:
            return False
        text = await asyncio.to_thread(extract_office_text, absolute_path(row.storage_path), ext)
        if not text:
            return False
        row.content_text = text
        await db.commit()
        return True


__all__ = [
    "DOCX_MIME",
    "OFFICE_ALLOWLIST",
    "OFFICE_EXTS",
    "ZIP_OFFICE_EXTS",
    "check_zip_container",
    "extract_and_store",
    "extract_office_text",
    "office_ext",
]
