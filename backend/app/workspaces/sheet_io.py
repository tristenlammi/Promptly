"""Spreadsheet import / export: .xlsx, .ods and .csv ↔ Fortune-sheet.

R-132 added values-and-formulas import/export; WP-4.4 carries the look of
a workbook both ways, so a budget comes in looking like the budget:

* fonts (bold, italic, underline, strike-through, size, family, colour),
  fills, alignment and wrapping, number formats (and the text they show);
* merged cells, column widths, row heights, hidden rows/columns/tabs,
  frozen panes, borders, web links;
* dates as real dates (an Excel serial with a date format), formulas with
  their cached results.

What Fortune-sheet can't hold (charts, pictures, conditional formatting,
data validation, comments, pivot tables, macros) is listed in ``notes`` so
the person knows what didn't come across.

Fortune-sheet stores a workbook as a list of tabs, each with sparse
``celldata``: ``[{"r": row, "c": col, "v": {"v": value, "m": shown,
"f": "=formula", "ct": {"fa": format, "t": type}, "bl": 1, …}}]`` plus a
``config`` (``merge``, ``columnlen``, ``rowlen``, ``borderInfo``, …).

Conversion runs here, not in the browser: the format spike found the
browser library (fortune-excel) crashing on its own imports and losing the
difference between text and numbers. openpyxl is already in the image and
reads through defusedxml.
"""
from __future__ import annotations

import colorsys
import csv
import io
import re
import uuid
from datetime import date, datetime, time
from typing import Any

from app.workspaces.sheet_numfmt import datetime_to_serial, format_value, general, is_date_format

# A phone-call-sized spreadsheet, not a data warehouse.
MAX_ROWS = 10_000
MAX_COLS = 200
MAX_CELLS = 400_000
# Above this a workbook is read values-only (openpyxl's full mode, which
# formatting needs, holds the whole file in memory).
FULL_MODE_MAX_BYTES = 8 * 1024 * 1024

SPREADSHEET_EXTS = (".xlsx", ".xlsm", ".ods", ".csv", ".tsv")


class SheetImportError(ValueError):
    pass


# --------------------------------------------------------------- helpers


def _cell_value(v: Any) -> tuple[Any, str | None]:
    """``(value, formula)`` of a Fortune cell's ``v``."""
    if isinstance(v, dict):
        f = v.get("f")
        val = v.get("v") if v.get("v") is not None else v.get("m")
        ct = v.get("ct") if isinstance(v.get("ct"), dict) else {}
        if val is None and ct.get("t") == "inlineStr" and isinstance(ct.get("s"), list):
            val = "".join(str(x.get("v", "")) for x in ct["s"] if isinstance(x, dict))
        return val, (str(f) if f else None)
    return v, None


def _tabs(data: Any) -> list[dict]:
    return [t for t in (data or []) if isinstance(t, dict)] if isinstance(data, list) else []


def _grid(tab: dict) -> dict[tuple[int, int], dict]:
    out: dict[tuple[int, int], dict] = {}
    for cell in tab.get("celldata") or []:
        try:
            out[(int(cell["r"]), int(cell["c"]))] = cell
        except (KeyError, TypeError, ValueError):
            continue
    if not out and isinstance(tab.get("data"), list):
        for r, row in enumerate(tab["data"]):
            for c, v in enumerate(row or []):
                if v is not None:
                    out[(r, c)] = {"r": r, "c": c, "v": v}
    return out


def _shown(v: Any) -> str:
    if isinstance(v, dict):
        if v.get("m") is not None:
            return str(v["m"])
        value, formula = _cell_value(v)
        if value is None:
            return formula or ""
        return general(value)
    return "" if v is None else general(v)


def _hex6(color: Any) -> str | None:
    """A Fortune colour (``#abc``, ``#aabbcc``, ``rgb(r, g, b)``) → ``AABBCC``."""
    if not isinstance(color, str):
        return None
    c = color.strip()
    m = re.fullmatch(r"#?([0-9a-fA-F]{6})", c)
    if m:
        return m.group(1).upper()
    m = re.fullmatch(r"#([0-9a-fA-F]{3})", c)
    if m:
        return "".join(ch * 2 for ch in m.group(1)).upper()
    m = re.fullmatch(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*(?:,\s*[\d.]+\s*)?\)", c)
    if m:
        return "".join(f"{min(255, int(x)):02X}" for x in m.groups())
    return None


def _num(v: Any) -> float | int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        try:
            f = float(v.replace(",", "")) if re.fullmatch(r"-?[\d,]*\.?\d+(?:[eE][+-]?\d+)?", v.strip()) else None
        except ValueError:
            return None
        if f is None:
            return None
        return int(f) if f.is_integer() and "." not in v and "e" not in v.lower() else f
    return None


def _fortune_cell(value: Any, formula: str | None = None, fmt: str | None = None) -> dict:
    """A Fortune cell for ``value`` (shown through Excel format ``fmt``)."""
    cell: dict[str, Any] = {}
    fa = fmt if fmt and fmt != "General" else "General"
    if isinstance(value, (datetime, date, time)):
        serial = datetime_to_serial(value)
        if fa == "General":
            fa = "hh:mm:ss" if isinstance(value, time) else (
                "yyyy-mm-dd hh:mm" if isinstance(value, datetime) and (value.hour or value.minute) else "yyyy-mm-dd"
            )
        cell["v"] = serial
        cell["ct"] = {"fa": fa, "t": "d"}
        shown = format_value(serial, fa)
    elif isinstance(value, bool):
        cell["v"] = "TRUE" if value else "FALSE"
        cell["ct"] = {"fa": "General", "t": "b"}
        shown = cell["v"]
    elif isinstance(value, (int, float)):
        cell["v"] = value
        cell["ct"] = {"fa": fa, "t": "d" if is_date_format(fa) else "n"}
        shown = format_value(value, fa)
    elif value is None and formula and fa not in ("General", "@"):
        # A formula saved without its result keeps its number format, so
        # the editor's recalculation shows "1,380.50", not "1380.5".
        cell["v"] = None
        cell["ct"] = {"fa": fa, "t": "d" if is_date_format(fa) else "n"}
        shown = None
    else:
        cell["v"] = value
        cell["ct"] = {"fa": "@", "t": "s"} if fa == "@" else {"fa": "General", "t": "g"}
        shown = "" if value is None else str(value)
    if shown is not None and not (formula and value is None):
        cell["m"] = shown  # absent → the editor formats / recalculates it
    if formula:
        cell["f"] = formula if formula.startswith("=") else "=" + formula
    return cell


# ---------------------------------------------------------- xlsx → Fortune

_DEFAULT_THEME = [
    "FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31",
    "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72",
]
_BORDER_STYLE = {
    "thin": 1, "hair": 2, "dotted": 3, "dashed": 4, "dashDot": 5, "dashDotDot": 6,
    "double": 7, "medium": 8, "mediumDashed": 9, "mediumDashDot": 10,
    "mediumDashDotDot": 11, "slantDashDot": 12, "thick": 13,
}
_BORDER_NAME = {v: k for k, v in _BORDER_STYLE.items()}
_H_ALIGN = {"center": 0, "centerContinuous": 0, "left": 1, "right": 2, "justify": 1, "distributed": 0}
_V_ALIGN = {"center": 0, "top": 1, "bottom": 2, "justify": 0, "distributed": 0}
_UNDERLINE = {"single": 1, "double": 2, "singleAccounting": 3, "doubleAccounting": 4}


def _theme_palette(wb: Any) -> list[str]:
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return list(_DEFAULT_THEME)
    try:
        from defusedxml import ElementTree as DET

        root = DET.fromstring(raw, forbid_dtd=True)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        scheme = root.find(f".//{a}clrScheme")
        found: dict[str, str] = {}
        for child in list(scheme) if scheme is not None else []:
            name = child.tag.rsplit("}", 1)[-1]
            inner = child[0] if len(child) else None
            if inner is None:
                continue
            val = inner.get("lastClr") or inner.get("val")
            if val and re.fullmatch(r"[0-9A-Fa-f]{6}", val):
                found[name] = val.upper()
        order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                 "accent4", "accent5", "accent6", "hlink", "folHlink"]
        return [found.get(n, d) for n, d in zip(order, _DEFAULT_THEME)]
    except Exception:  # noqa: BLE001 — a malformed theme just means defaults
        return list(_DEFAULT_THEME)


def _tint(hex6: str, tint: float) -> str:
    if not tint:
        return hex6
    r, g, b = (int(hex6[i : i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _xl_color(color: Any, palette: list[str]) -> str | None:
    """An openpyxl ``Color`` → ``#RRGGBB`` (theme colours and tints too)."""
    if color is None:
        return None
    try:
        kind = getattr(color, "type", None)
        if kind == "rgb":
            rgb = color.rgb
            if not isinstance(rgb, str) or len(rgb) < 6:
                return None
            if len(rgb) == 8 and rgb[:2] == "00" and rgb[2:] == "000000":
                return None  # "auto" written as transparent black
            return "#" + rgb[-6:].upper()
        if kind == "theme":
            idx = int(color.theme)
            if 0 <= idx < len(palette):
                return "#" + _tint(palette[idx], float(color.tint or 0))
            return None
        if kind == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX

            idx = int(color.indexed)
            if 0 <= idx < len(COLOR_INDEX) and idx < 64:
                return "#" + COLOR_INDEX[idx][-6:].upper()
    except Exception:  # noqa: BLE001
        return None
    return None


def _align_numbers(v: dict) -> dict:
    """Excel and LibreOffice right-align numbers and dates unless told
    otherwise; Fortune-sheet left-aligns everything. Say so explicitly."""
    if "ht" not in v and isinstance(v.get("ct"), dict) and v["ct"].get("t") in ("n", "d"):
        v["ht"] = 2
    return v


def _style_of(cell: Any, palette: list[str]) -> dict:
    """Fortune style keys for an openpyxl cell (only what differs from plain)."""
    out: dict[str, Any] = {}
    font = cell.font
    if font is not None:
        if font.b:
            out["bl"] = 1
        if font.i:
            out["it"] = 1
        if font.strike:
            out["cl"] = 1
        if font.u:
            out["un"] = _UNDERLINE.get(font.u, 1)
        if font.sz:
            try:
                size = float(font.sz)
                out["fs"] = int(size) if size.is_integer() else size
            except (TypeError, ValueError):
                pass
        if font.name:
            out["ff"] = str(font.name)
        fc = _xl_color(font.color, palette)
        if fc and fc != "#000000":
            out["fc"] = fc
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = _xl_color(fill.fgColor, palette) or _xl_color(fill.start_color, palette)
        if bg:
            out["bg"] = bg
    al = cell.alignment
    if al is not None:
        if al.horizontal in _H_ALIGN:
            out["ht"] = _H_ALIGN[al.horizontal]
        if al.vertical in _V_ALIGN:
            out["vt"] = _V_ALIGN[al.vertical]
        if al.wrap_text:
            out["tb"] = "2"
    return out


def _border_of(cell: Any, palette: list[str]) -> dict | None:
    b = cell.border
    if b is None:
        return None
    sides = {}
    for key, side in (("l", b.left), ("r", b.right), ("t", b.top), ("b", b.bottom)):
        if side is not None and side.style in _BORDER_STYLE:
            sides[key] = {"style": _BORDER_STYLE[side.style], "color": _xl_color(side.color, palette) or "#000000"}
    return sides or None


def _xlsx_notes(wb: Any) -> list[str]:
    seen: dict[str, int] = {}

    def add(key: str, n: int = 1) -> None:
        seen[key] = seen.get(key, 0) + n

    for ws in wb.worksheets:
        if getattr(ws, "_charts", None):
            add("chart", len(ws._charts))
        if getattr(ws, "_images", None):
            add("image", len(ws._images))
        try:
            if len(ws.conditional_formatting):
                add("conditional")
        except Exception:  # noqa: BLE001
            pass
        dv = getattr(getattr(ws, "data_validations", None), "dataValidation", None)
        if dv:
            add("validation")
        if getattr(ws, "_pivots", None):
            add("pivot", len(ws._pivots))
        comments = sum(1 for cell in getattr(ws, "_cells", {}).values() if getattr(cell, "comment", None))
        if comments:
            add("comment", comments)
    notes = []
    if seen.get("chart"):
        n = seen["chart"]
        notes.append(f"{n} chart{'s' if n != 1 else ''} weren't brought over." if n != 1 else "1 chart wasn't brought over.")
    if seen.get("image"):
        n = seen["image"]
        notes.append("1 picture wasn't brought over." if n == 1 else f"{n} pictures weren't brought over.")
    if seen.get("comment"):
        n = seen["comment"]
        notes.append("1 cell comment wasn't brought over." if n == 1 else f"{n} cell comments weren't brought over.")
    if seen.get("conditional"):
        notes.append("Conditional formatting isn't kept (the cells keep their values).")
    if seen.get("validation"):
        notes.append("Drop-down lists and other data validation aren't kept.")
    if seen.get("pivot"):
        notes.append("Pivot tables come in as plain values.")
    return notes


def _xlsx_full(raw: bytes) -> tuple[list[dict], list[str]]:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import column_index_from_string

    wb = load_workbook(io.BytesIO(raw), data_only=False)
    # Formula results live in the cached values, read separately (streamed).
    cached: dict[str, dict[tuple[int, int], Any]] = {}
    formula_at: dict[str, set[tuple[int, int]]] = {}
    for ws in wb.worksheets:
        pos = {
            k for k, cell in ws._cells.items()
            if isinstance(cell.value, str) and cell.value.startswith("=")
            or getattr(cell, "data_type", None) == "f"
        }
        if pos:
            formula_at[ws.title] = pos
    if formula_at:
        wv = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        try:
            for title, pos in formula_at.items():
                want = cached.setdefault(title, {})
                max_r = min(MAX_ROWS, max(r for r, _ in pos))
                for row in wv[title].iter_rows(min_row=1, max_row=max_r):
                    for c in row:
                        key = (getattr(c, "row", None), getattr(c, "column", None))
                        if key in pos:
                            want[key] = c.value
        finally:
            wv.close()

    palette = _theme_palette(wb)
    notes = _xlsx_notes(wb)
    tabs: list[dict] = []
    total = 0
    truncated = False
    for order, ws in enumerate(wb.worksheets):
        sid = uuid.uuid4().hex[:12]
        celldata: list[dict] = []
        config: dict[str, Any] = {}
        borders: list[dict] = []
        hyperlinks: dict[str, dict] = {}
        merge: dict[str, dict] = {}
        covered: dict[tuple[int, int], tuple[int, int]] = {}
        for rng in ws.merged_cells.ranges:
            r0, c0 = rng.min_row - 1, rng.min_col - 1
            if r0 >= MAX_ROWS or c0 >= MAX_COLS:
                continue
            rs = min(rng.max_row, MAX_ROWS) - rng.min_row + 1
            cs = min(rng.max_col, MAX_COLS) - rng.min_col + 1
            merge[f"{r0}_{c0}"] = {"r": r0, "c": c0, "rs": rs, "cs": cs}
            for rr in range(r0, r0 + rs):
                for cc in range(c0, c0 + cs):
                    covered[(rr, cc)] = (r0, c0)
        cells_by_pos: dict[tuple[int, int], dict] = {}
        for (row, col), cell in sorted(ws._cells.items()):
            r, c = row - 1, col - 1
            if r >= MAX_ROWS or c >= MAX_COLS:
                truncated = True
                continue
            border = _border_of(cell, palette)
            if border:
                borders.append({"rangeType": "cell", "value": {"row_index": r, "col_index": c, **border}})
            if isinstance(cell, MergedCell):
                continue
            style = _style_of(cell, palette)
            raw_value = cell.value
            formula = None
            if isinstance(raw_value, str) and raw_value.startswith("="):
                formula = raw_value
                raw_value = cached.get(ws.title, {}).get((row, col))
            elif getattr(raw_value, "text", None) is not None and type(raw_value).__name__ == "ArrayFormula":
                formula = str(raw_value.text)
                raw_value = cached.get(ws.title, {}).get((row, col))
            if raw_value is None and formula is None:
                if "bg" not in style:
                    continue  # nothing to show
                v: dict[str, Any] = {}
            else:
                if type(raw_value).__name__ == "CellRichText":
                    raw_value = str(raw_value)
                v = _fortune_cell(raw_value, formula, cell.number_format)
            v.update(style)
            link = getattr(cell, "hyperlink", None)
            target = getattr(link, "target", None) if link is not None else None
            if target and re.match(r"(?i)^(https?:|mailto:)", str(target)):
                hyperlinks[f"{r}_{c}"] = {"linkType": "webpage", "linkAddress": str(target)}
                v["hl"] = {"r": r, "c": c, "id": sid}
            total += 1
            if total > MAX_CELLS:
                raise SheetImportError(f"That workbook is too big to import (more than {MAX_CELLS:,} cells).")
            cells_by_pos[(r, c)] = _align_numbers(v)
        for (r, c), (r0, c0) in covered.items():
            v = cells_by_pos.setdefault((r, c), {})
            if (r, c) == (r0, c0):
                m = merge[f"{r0}_{c0}"]
                v["mc"] = dict(m)
            else:
                v["mc"] = {"r": r0, "c": c0}
        for (r, c), v in sorted(cells_by_pos.items()):
            celldata.append({"r": r, "c": c, "v": v})

        columnlen: dict[str, int] = {}
        colhidden: dict[str, int] = {}
        for key, dim in ws.column_dimensions.items():
            try:
                lo = dim.min or column_index_from_string(key)
                hi = dim.max or lo
            except (TypeError, ValueError):
                continue
            for col in range(lo, min(hi, MAX_COLS) + 1):
                if dim.width and getattr(dim, "customWidth", True):
                    columnlen[str(col - 1)] = max(1, round(float(dim.width) * 7 + 5))
                if dim.hidden:
                    colhidden[str(col - 1)] = 0
        rowlen: dict[str, int] = {}
        rowhidden: dict[str, int] = {}
        for idx, dim in ws.row_dimensions.items():
            if not isinstance(idx, int) or idx > MAX_ROWS:
                continue
            if dim.height:
                rowlen[str(idx - 1)] = max(1, round(float(dim.height) * 4 / 3))
            if dim.hidden:
                rowhidden[str(idx - 1)] = 0
        if merge:
            config["merge"] = merge
        if columnlen:
            config["columnlen"] = columnlen
        if rowlen:
            config["rowlen"] = rowlen
        if colhidden:
            config["colhidden"] = colhidden
        if rowhidden:
            config["rowhidden"] = rowhidden
        if borders:
            config["borderInfo"] = borders
        tab: dict[str, Any] = {
            "id": sid,
            "name": ws.title[:60] or f"Sheet{order + 1}",
            "order": order,
            "celldata": celldata,
            "config": config,
        }
        if hyperlinks:
            tab["hyperlink"] = hyperlinks
        if ws.sheet_state != "visible":
            tab["hide"] = 1
        frozen = _frozen_from(ws.freeze_panes)
        if frozen:
            tab["frozen"] = frozen
        tabs.append(tab)
    if truncated:
        notes.append(f"Only the first {MAX_ROWS:,} rows and {MAX_COLS} columns of each tab came across.")
    if tabs and all(t.get("hide") for t in tabs):
        tabs[0].pop("hide", None)
    return tabs, notes


def _frozen_from(pane: Any) -> dict | None:
    if not pane or not isinstance(pane, str):
        return None
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    try:
        letters, row = coordinate_from_string(pane)
        col = column_index_from_string(letters)
    except (TypeError, ValueError):
        return None
    rows, cols = row - 1, col - 1
    if rows > 0 and cols > 0:
        return {"type": "rangeBoth", "range": {"row_focus": rows - 1, "column_focus": cols - 1}}
    if rows > 0:
        return {"type": "rangeRow", "range": {"row_focus": rows - 1, "column_focus": 0}}
    if cols > 0:
        return {"type": "rangeColumn", "range": {"row_focus": 0, "column_focus": cols - 1}}
    return None


def _xlsx_values_only(raw: bytes) -> list[dict]:
    """The R-132 reader: values and formulas, streamed (big files)."""
    from openpyxl import load_workbook

    wb_f = load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
    wb_v = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    tabs: list[dict] = []
    total = 0
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            celldata = []
            for row_f, row_v in zip(
                ws_f.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS),
                ws_v.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS),
            ):
                for cf, cv in zip(row_f, row_v):
                    raw_value = getattr(cf, "value", None)
                    if raw_value is None:
                        continue
                    formula = raw_value if isinstance(raw_value, str) and raw_value.startswith("=") else None
                    value = cv.value if formula else raw_value
                    fmt = getattr(cf, "number_format", None)
                    celldata.append(
                        {"r": cf.row - 1, "c": cf.column - 1, "v": _align_numbers(_fortune_cell(value, formula, fmt))}
                    )
                    total += 1
                    if total > MAX_CELLS:
                        raise SheetImportError(
                            f"That workbook is too big to import (more than {MAX_CELLS:,} cells)."
                        )
            tabs.append({"name": ws_f.title[:60] or f"Sheet{len(tabs) + 1}", "celldata": celldata})
    finally:
        wb_f.close()
        wb_v.close()
    return tabs


def from_xlsx_with_notes(raw: bytes) -> tuple[list[dict], list[str]]:
    """``(tabs, notes)`` for an .xlsx / .xlsm — formatted when it's a
    normal-sized file, values-only (with a note) when it's huge."""
    try:
        if len(raw) <= FULL_MODE_MAX_BYTES:
            tabs, notes = _xlsx_full(raw)
        else:
            tabs = _xlsx_values_only(raw)
            notes = ["This workbook is large, so only its values and formulas came across (no formatting)."]
    except SheetImportError:
        raise
    except Exception as exc:  # noqa: BLE001 — openpyxl raises a zoo
        raise SheetImportError("That file isn't a readable .xlsx workbook.") from exc
    return (tabs or [{"name": "Sheet1", "celldata": []}]), notes


def from_xlsx(raw: bytes) -> list[dict]:
    return from_xlsx_with_notes(raw)[0]


# ---------------------------------------------------------- ods → Fortune

_ODF = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
}


def _q(prefix: str, name: str) -> str:
    return "{%s}%s" % (_ODF[prefix], name)


def _length_px(value: str | None) -> int | None:
    m = re.fullmatch(r"\s*([\d.]+)\s*(in|cm|mm|pt|px)\s*", value or "")
    if not m:
        return None
    n = float(m.group(1))
    factor = {"in": 96, "cm": 96 / 2.54, "mm": 96 / 25.4, "pt": 96 / 72, "px": 1}[m.group(2)]
    return max(1, round(n * factor))


def _of_to_excel(formula: str) -> str:
    """OpenFormula (``of:=SUM([.B3:.B4])``) → Excel (``=SUM(B3:B4)``)."""
    f = formula
    for prefix in ("of:", "oooc:", "msoxl:"):
        if f.startswith(prefix):
            f = f[len(prefix) :]
            break
    if not f.startswith("="):
        f = "=" + f

    def ref(m: re.Match) -> str:
        parts = []
        sheet_seen = None
        for part in m.group(1).split(":"):
            part = part.strip()
            if part.startswith("."):
                cell = part[1:]
                parts.append(cell if sheet_seen is None or parts else f"{sheet_seen}!{cell}")
            else:
                sheet, _, cell = part.rpartition(".")
                sheet = sheet.lstrip("$")
                if sheet and not re.fullmatch(r"[A-Za-z_][\w]*", sheet.strip("'")):
                    sheet = "'" + sheet.strip("'") + "'"
                sheet_seen = sheet
                parts.append(f"{sheet}!{cell}" if sheet else cell)
        return ":".join(parts)

    out, quoted = [], False
    f = re.sub(r"\[([^\]]+)\]", ref, f)
    for ch in f:
        if ch == '"':
            quoted = not quoted
        out.append("," if ch == ";" and not quoted else ch)
    return "".join(out)


def _ods_styles(root: Any) -> tuple[dict[str, dict], dict[str, int], dict[str, int]]:
    cell_styles: dict[str, dict] = {}
    col_widths: dict[str, int] = {}
    row_heights: dict[str, int] = {}
    auto = root.find(_q("office", "automatic-styles"))
    for st in list(auto) if auto is not None else []:
        if st.tag != _q("style", "style"):
            continue
        name = st.get(_q("style", "name"))
        family = st.get(_q("style", "family"))
        if family == "table-column":
            props = st.find(_q("style", "table-column-properties"))
            px = _length_px(props.get(_q("style", "column-width")) if props is not None else None)
            if px:
                col_widths[name] = px
        elif family == "table-row":
            props = st.find(_q("style", "table-row-properties"))
            px = _length_px(props.get(_q("style", "row-height")) if props is not None else None)
            if px:
                row_heights[name] = px
        elif family == "table-cell":
            out: dict[str, Any] = {}
            text = st.find(_q("style", "text-properties"))
            if text is not None:
                if text.get(_q("fo", "font-weight")) == "bold":
                    out["bl"] = 1
                if text.get(_q("fo", "font-style")) == "italic":
                    out["it"] = 1
                color = _hex6(text.get(_q("fo", "color")))
                if color and color != "000000":
                    out["fc"] = "#" + color
                size = re.fullmatch(r"([\d.]+)pt", text.get(_q("fo", "font-size")) or "")
                if size:
                    s = float(size.group(1))
                    out["fs"] = int(s) if s.is_integer() else s
                if text.get(_q("style", "font-name")):
                    out["ff"] = text.get(_q("style", "font-name"))
                if (text.get(_q("style", "text-underline-style")) or "none") != "none":
                    out["un"] = 1
                if (text.get(_q("style", "text-line-through-style")) or "none") != "none":
                    out["cl"] = 1
            cellp = st.find(_q("style", "table-cell-properties"))
            if cellp is not None:
                bg = _hex6(cellp.get(_q("fo", "background-color")))
                if bg:
                    out["bg"] = "#" + bg
                if cellp.get(_q("fo", "wrap-option")) == "wrap":
                    out["tb"] = "2"
                va = cellp.get(_q("style", "vertical-align"))
                if va in ("top", "middle", "bottom"):
                    out["vt"] = {"middle": 0, "top": 1, "bottom": 2}[va]
            para = st.find(_q("style", "paragraph-properties"))
            if para is not None:
                ta = para.get(_q("fo", "text-align"))
                if ta in ("center",):
                    out["ht"] = 0
                elif ta in ("start", "left"):
                    out["ht"] = 1
                elif ta in ("end", "right"):
                    out["ht"] = 2
            cell_styles[name] = out
    return cell_styles, col_widths, row_heights


def _ods_value(el: Any) -> tuple[Any, str | None, str | None]:
    """``(value, Excel format, shown text)`` of an ODS cell."""
    office = lambda n: el.get(_q("office", n))  # noqa: E731
    kind = office("value-type")
    shown = "\n".join("".join(p.itertext()) for p in el.findall(_q("text", "p"))) or None
    if kind in ("float", "percentage", "currency"):
        try:
            n = float(office("value"))
        except (TypeError, ValueError):
            return shown, None, shown
        n = int(n) if n.is_integer() else n
        fmt = "0.00%" if kind == "percentage" else ("#,##0.00" if kind == "currency" else None)
        return n, fmt, shown
    if kind == "date":
        raw = office("date-value") or ""
        try:
            dt = datetime.fromisoformat(raw)
            fmt = "yyyy-mm-dd hh:mm" if "T" in raw else "yyyy-mm-dd"
            return datetime_to_serial(dt), fmt, shown
        except ValueError:
            return shown, None, shown
    if kind == "time":
        m = re.fullmatch(r"PT(?:(\d+)H)?(?:(\d+)M)?(?:([\d.]+)S)?", office("time-value") or "")
        if m:
            h, mi, s = (float(x or 0) for x in m.groups())
            return (h * 3600 + mi * 60 + s) / 86400, "hh:mm", shown
        return shown, None, shown
    if kind == "boolean":
        return (office("boolean-value") or "").lower() == "true", None, shown
    return shown, None, shown


def from_ods_with_notes(raw: bytes) -> tuple[list[dict], list[str]]:
    import zipfile

    from defusedxml import ElementTree as DET

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            content = zf.read("content.xml")
        root = DET.fromstring(content, forbid_dtd=True)
    except Exception as exc:  # noqa: BLE001
        raise SheetImportError("That file isn't a readable .ods spreadsheet.") from exc
    cell_styles, col_widths, row_heights = _ods_styles(root)
    body = root.find(_q("office", "body"))
    sheet_root = body.find(_q("office", "spreadsheet")) if body is not None else None
    tabs: list[dict] = []
    notes: list[str] = []
    total = 0
    truncated = False
    charts = 0
    for order, table in enumerate(sheet_root.findall(_q("table", "table")) if sheet_root is not None else []):
        celldata: list[dict] = []
        merge: dict[str, dict] = {}
        columnlen: dict[str, int] = {}
        rowlen: dict[str, int] = {}
        c = 0
        for col in table.iter(_q("table", "table-column")):
            rep = int(col.get(_q("table", "number-columns-repeated")) or 1)
            width = col_widths.get(col.get(_q("table", "style-name")) or "")
            if width:
                for cc in range(c, min(c + rep, MAX_COLS)):
                    columnlen[str(cc)] = width
            c += rep
            if c >= MAX_COLS:
                break
        charts += len(list(table.iter("{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}frame")))
        r = 0
        for row in table.iter(_q("table", "table-row")):
            if r >= MAX_ROWS:
                truncated = True
                break
            rrep = int(row.get(_q("table", "number-rows-repeated")) or 1)
            height = row_heights.get(row.get(_q("table", "style-name")) or "")
            cells: list[tuple[int, dict]] = []
            c = 0
            for el in row:
                if el.tag not in (_q("table", "table-cell"), _q("table", "covered-table-cell")):
                    continue
                crep = int(el.get(_q("table", "number-columns-repeated")) or 1)
                if el.tag == _q("table", "table-cell"):
                    value, fmt, shown = _ods_value(el)
                    formula = el.get(_q("table", "formula"))
                    style = cell_styles.get(el.get(_q("table", "style-name")) or "", {})
                    if value is not None or formula or "bg" in style:
                        if value is None and not formula:
                            v: dict[str, Any] = {}
                        else:
                            v = _fortune_cell(value, _of_to_excel(formula) if formula else None, fmt)
                            if shown is not None and v.get("ct", {}).get("t") in ("n", "d"):
                                v["m"] = shown  # LibreOffice's own rendering
                        v.update(style)
                        _align_numbers(v)
                        for k in range(min(crep, MAX_COLS - c)):
                            cells.append((c + k, dict(v)))
                    cs = int(el.get(_q("table", "number-columns-spanned")) or 1)
                    rs = int(el.get(_q("table", "number-rows-spanned")) or 1)
                    if (cs > 1 or rs > 1) and c < MAX_COLS:
                        merge[f"{r}_{c}"] = {"r": r, "c": c, "rs": min(rs, MAX_ROWS - r), "cs": min(cs, MAX_COLS - c)}
                elif c >= MAX_COLS:
                    pass
                c += crep
                if c >= MAX_COLS:
                    if any(True for _ in el.itertext()):
                        truncated = True
                    break
            for k in range(min(rrep, MAX_ROWS - r) if cells or height else 0):
                if height:
                    rowlen[str(r + k)] = height
                for cc, v in cells:
                    celldata.append({"r": r + k, "c": cc, "v": dict(v)})
                    total += 1
                    if total > MAX_CELLS:
                        raise SheetImportError(f"That spreadsheet is too big to import (more than {MAX_CELLS:,} cells).")
            r += rrep
        by_pos = {(cell["r"], cell["c"]): cell for cell in celldata}
        for key, m in merge.items():
            for rr in range(m["r"], m["r"] + m["rs"]):
                for cc in range(m["c"], m["c"] + m["cs"]):
                    cell = by_pos.get((rr, cc))
                    if cell is None:
                        cell = {"r": rr, "c": cc, "v": {}}
                        celldata.append(cell)
                        by_pos[(rr, cc)] = cell
                    cell["v"]["mc"] = dict(m) if (rr, cc) == (m["r"], m["c"]) else {"r": m["r"], "c": m["c"]}
        celldata.sort(key=lambda x: (x["r"], x["c"]))
        config: dict[str, Any] = {}
        if merge:
            config["merge"] = merge
        if columnlen:
            config["columnlen"] = columnlen
        if rowlen:
            config["rowlen"] = rowlen
        tabs.append({
            "id": uuid.uuid4().hex[:12],
            "name": (table.get(_q("table", "name")) or f"Sheet{order + 1}")[:60],
            "order": order,
            "celldata": celldata,
            "config": config,
        })
    if charts:
        notes.append("Charts and pictures weren't brought over.")
    if truncated:
        notes.append(f"Only the first {MAX_ROWS:,} rows and {MAX_COLS} columns of each tab came across.")
    return (tabs or [{"name": "Sheet1", "celldata": []}]), notes


# ---------------------------------------------------------- csv → Fortune


def ensure_tab_ids(data: Any) -> bool:
    """Give every tab a stable id, in place; True when one was missing.

    Tabs are shared live by id (WP-6.3). A tab saved without one got a
    different random id in each browser until somebody saved — now the
    server names it once."""
    changed = False
    for tab in _tabs(data):
        if not tab.get("id"):
            tab["id"] = uuid.uuid4().hex[:12]
            changed = True
    return changed


def from_csv(raw: bytes, name: str = "Sheet1") -> list[dict]:
    text = raw.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    celldata = []
    for r, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if r >= MAX_ROWS:
            break
        for c, val in enumerate(row[:MAX_COLS]):
            if val == "":
                continue
            value: Any = val
            # Numbers become numbers — but not codes with leading zeros.
            if not (len(val) > 1 and val.startswith("0") and not val.startswith("0.")):
                try:
                    value = int(val) if val.lstrip("-").isdigit() else float(val)
                except ValueError:
                    value = val
            celldata.append({"r": r, "c": c, "v": _align_numbers(_fortune_cell(value))})
            if len(celldata) > MAX_CELLS:
                raise SheetImportError("That file is too big to import.")
    return [{"id": uuid.uuid4().hex[:12], "name": name[:60] or "Sheet1", "celldata": celldata}]


def import_spreadsheet(raw: bytes, filename: str) -> tuple[list[dict], list[str]]:
    """Any supported spreadsheet → ``(tabs, notes)``. Zip formats pass the
    zip-bomb guard first. Raises :class:`SheetImportError`."""
    import os

    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm", ".ods"):
        from app.files.office_formats import check_zip_container
        from app.files.safety import UnsafeUploadError

        try:
            check_zip_container(io.BytesIO(raw), ".xlsx" if ext == ".xlsm" else ext)  # type: ignore[arg-type]
        except UnsafeUploadError as exc:
            raise SheetImportError(str(exc)) from exc
    if ext in (".xlsx", ".xlsm"):
        tabs, notes = from_xlsx_with_notes(raw)
        if ext == ".xlsm":
            notes.append("Macros aren't kept.")
        return tabs, notes
    if ext == ".ods":
        return from_ods_with_notes(raw)
    if ext in (".csv", ".tsv", ".txt"):
        return from_csv(raw, os.path.splitext(os.path.basename(filename))[0] or "Sheet1"), []
    raise SheetImportError("Choose an .xlsx, .ods or .csv file.")


# ---------------------------------------------------------- Fortune → xlsx


def _frozen_cell(frozen: Any) -> str | None:
    from openpyxl.utils import get_column_letter

    if not isinstance(frozen, dict):
        return None
    kind = frozen.get("type")
    rng = frozen.get("range") if isinstance(frozen.get("range"), dict) else {}
    rf = int(rng.get("row_focus") or 0)
    cf = int(rng.get("column_focus") or 0)
    if kind == "row":
        return "A2"
    if kind == "column":
        return "B1"
    if kind == "both":
        return "B2"
    if kind == "rangeRow":
        return f"A{rf + 2}"
    if kind == "rangeColumn":
        return f"{get_column_letter(cf + 2)}1"
    if kind == "rangeBoth":
        return f"{get_column_letter(cf + 2)}{rf + 2}"
    return None


def _cell_borders(tab: dict) -> dict[tuple[int, int], dict[str, tuple[int, str]]]:
    """Fortune ``borderInfo`` (cell and range forms) → per-cell sides."""
    out: dict[tuple[int, int], dict[str, tuple[int, str]]] = {}
    info = (tab.get("config") or {}).get("borderInfo") if isinstance(tab.get("config"), dict) else None

    def put(r: int, c: int, side: str, style: Any, color: Any) -> None:
        if r < 0 or c < 0 or r >= MAX_ROWS or c >= MAX_COLS:
            return
        try:
            s = int(style)
        except (TypeError, ValueError):
            return
        if s in _BORDER_NAME:
            out.setdefault((r, c), {})[side] = (s, _hex6(color) or "000000")
        elif s == 0:
            out.get((r, c), {}).pop(side, None)

    for item in info if isinstance(info, list) else []:
        if not isinstance(item, dict):
            continue
        if item.get("rangeType") == "cell":
            val = item.get("value") or {}
            try:
                r, c = int(val["row_index"]), int(val["col_index"])
            except (KeyError, TypeError, ValueError):
                continue
            for side in ("l", "r", "t", "b"):
                if isinstance(val.get(side), dict):
                    put(r, c, side, val[side].get("style"), val[side].get("color"))
            continue
        kind, style, color = item.get("borderType"), item.get("style"), item.get("color")
        for rng in item.get("range") or []:
            try:
                r1, r2 = (int(x) for x in rng["row"])
                c1, c2 = (int(x) for x in rng["column"])
            except (KeyError, TypeError, ValueError):
                continue
            r2, c2 = min(r2, MAX_ROWS - 1), min(c2, MAX_COLS - 1)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    if kind == "border-none":
                        out.pop((r, c), None)
                        continue
                    sides = {
                        "border-all": "lrtb",
                        "border-outside": ("l" if c == c1 else "") + ("r" if c == c2 else "")
                        + ("t" if r == r1 else "") + ("b" if r == r2 else ""),
                        "border-inside": ("r" if c < c2 else "") + ("b" if r < r2 else ""),
                        "border-horizontal": "b" if r < r2 else "",
                        "border-vertical": "r" if c < c2 else "",
                        "border-left": "l" if c == c1 else "",
                        "border-right": "r" if c == c2 else "",
                        "border-top": "t" if r == r1 else "",
                        "border-bottom": "b" if r == r2 else "",
                    }.get(kind, "")
                    for side in sides:
                        put(r, c, side, style, color)
    return out


def _xl_value(v: Any) -> tuple[Any, str | None]:
    """``(cell value, number format)`` to write for a Fortune cell."""
    if not isinstance(v, dict):
        return v, None
    ct = v.get("ct") if isinstance(v.get("ct"), dict) else {}
    t, fa = ct.get("t"), ct.get("fa")
    value, formula = _cell_value(v)
    fmt = fa if isinstance(fa, str) and fa and fa != "General" else None
    if formula:
        return (formula if formula.startswith("=") else "=" + formula), fmt
    if t in ("n", "d"):
        n = _num(value)
        if n is not None:
            return n, fmt or ("yyyy-mm-dd" if t == "d" else None)
        return value, None
    if t == "b":
        return str(value).upper() in ("TRUE", "1"), None
    if isinstance(value, str) and fmt is None and t is None:
        n = _num(value)
        if n is not None and not (len(value) > 1 and value.startswith("0") and not value.startswith("0.")):
            return n, None
    if value == "":
        return None, fmt
    return value, ("@" if fa == "@" else fmt)


def to_xlsx(data: Any) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    wb.remove(wb.active)
    # No cached formula results are written: have Excel / LibreOffice work
    # them out when the file opens.
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    used: set[str] = set()
    tabs = _tabs(data) or [{"name": "Sheet1"}]
    if all(isinstance(t.get("order"), int) for t in tabs):
        tabs = sorted(tabs, key=lambda t: t["order"])
    h_align = {0: "center", 1: "left", 2: "right"}
    v_align = {0: "center", 1: "top", 2: "bottom"}
    u_style = {1: "single", 2: "double", 3: "singleAccounting", 4: "doubleAccounting"}
    for i, tab in enumerate(tabs):
        name = str(tab.get("name") or f"Sheet{i + 1}")[:31] or f"Sheet{i + 1}"
        for ch in '[]:*?/\\':
            name = name.replace(ch, "_")
        base, n = name, 2
        while name in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(name)
        config = tab.get("config") if isinstance(tab.get("config"), dict) else {}
        links = tab.get("hyperlink") if isinstance(tab.get("hyperlink"), dict) else {}
        for (r, c), cell in _grid(tab).items():
            if r >= MAX_ROWS or c >= MAX_COLS:
                continue
            v = cell.get("v")
            value, fmt = _xl_value(v)
            target = None
            if value is not None or fmt or isinstance(v, dict):
                target = ws.cell(row=r + 1, column=c + 1)
                if value is not None:
                    target.value = value
                if fmt:
                    target.number_format = fmt
            if not isinstance(v, dict) or target is None:
                continue
            if any(k in v for k in ("bl", "it", "fs", "ff", "fc", "un", "cl")):
                fc = _hex6(v.get("fc"))
                try:
                    size = float(v["fs"]) if v.get("fs") not in (None, "") else None
                except (TypeError, ValueError):
                    size = None
                target.font = Font(
                    name=v["ff"] if isinstance(v.get("ff"), str) and v["ff"] else None,
                    size=size,
                    bold=str(v.get("bl")) == "1" or None,
                    italic=str(v.get("it")) == "1" or None,
                    underline=u_style.get(int(v["un"])) if str(v.get("un") or "0").isdigit() and int(v.get("un") or 0) else None,
                    strike=str(v.get("cl")) == "1" or None,
                    color=("FF" + fc) if fc else None,
                )
            bg = _hex6(v.get("bg"))
            if bg:
                target.fill = PatternFill("solid", fgColor="FF" + bg)
            ht = v.get("ht")
            vt = v.get("vt")
            wrap = str(v.get("tb")) == "2"
            if ht is not None or vt is not None or wrap:
                try:
                    target.alignment = Alignment(
                        horizontal=h_align.get(int(ht)) if ht is not None and str(ht).isdigit() else None,
                        vertical=v_align.get(int(vt)) if vt is not None and str(vt).isdigit() else None,
                        wrap_text=wrap or None,
                    )
                except (TypeError, ValueError):
                    pass
            link = links.get(f"{r}_{c}") if isinstance(v.get("hl"), dict) else None
            if isinstance(link, dict) and link.get("linkType") == "webpage":
                addr = str(link.get("linkAddress") or "")
                if re.match(r"(?i)^(https?:|mailto:)", addr):
                    target.hyperlink = addr
        for (r, c), sides in _cell_borders(tab).items():
            kw = {
                {"l": "left", "r": "right", "t": "top", "b": "bottom"}[k]: Side(style=_BORDER_NAME[s], color="FF" + col)
                for k, (s, col) in sides.items()
            }
            if kw:
                ws.cell(row=r + 1, column=c + 1).border = Border(**kw)
        merges = config.get("merge") if isinstance(config.get("merge"), dict) else {}
        for m in merges.values():
            try:
                r, c, rs, cs = int(m["r"]), int(m["c"]), int(m.get("rs", 1)), int(m.get("cs", 1))
            except (KeyError, TypeError, ValueError):
                continue
            if rs * cs > 1 and r + rs <= MAX_ROWS and c + cs <= MAX_COLS:
                ws.merge_cells(start_row=r + 1, start_column=c + 1, end_row=r + rs, end_column=c + cs)
        for key, px in (config.get("columnlen") or {}).items():
            try:
                col = int(key)
                ws.column_dimensions[get_column_letter(col + 1)].width = round(max(0.0, (float(px) - 5) / 7), 2)
            except (TypeError, ValueError):
                continue
        for key, px in (config.get("rowlen") or {}).items():
            try:
                ws.row_dimensions[int(key) + 1].height = round(float(px) * 0.75, 2)
            except (TypeError, ValueError):
                continue
        for key in (config.get("colhidden") or {}):
            try:
                ws.column_dimensions[get_column_letter(int(key) + 1)].hidden = True
            except (TypeError, ValueError):
                continue
        for key in (config.get("rowhidden") or {}):
            try:
                ws.row_dimensions[int(key) + 1].hidden = True
            except (TypeError, ValueError):
                continue
        pane = _frozen_cell(tab.get("frozen"))
        if pane:
            ws.freeze_panes = pane
        if str(tab.get("hide")) == "1":
            ws.sheet_state = "hidden"
    if all(ws.sheet_state != "visible" for ws in wb.worksheets):
        wb.worksheets[0].sheet_state = "visible"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(data: Any, tab_index: int = 0) -> str:
    tabs = _tabs(data)
    tab = tabs[tab_index] if 0 <= tab_index < len(tabs) else (tabs[0] if tabs else {})
    grid = _grid(tab)
    buf = io.StringIO()
    if not grid:
        return ""
    max_r = max(r for r, _ in grid)
    max_c = max(c for _, c in grid)
    w = csv.writer(buf)
    for r in range(max_r + 1):
        row = []
        for c in range(max_c + 1):
            cell = grid.get((r, c))
            row.append("" if cell is None else _shown(cell.get("v")))
        w.writerow(row)
    return buf.getvalue()


# ---------------------------------------------------------- views


def flatten_text(data: Any, limit: int = 200_000) -> str:
    """Searchable text of a workbook: one line per row, cells joined by
    `` | ``, each tab under a ``## name`` heading."""
    parts: list[str] = []
    size = 0
    for i, tab in enumerate(_tabs(data)):
        rows: dict[int, list[tuple[int, str]]] = {}
        for (r, c), cell in _grid(tab).items():
            text = _shown(cell.get("v")).strip()
            if text:
                rows.setdefault(r, []).append((c, text))
        if not rows:
            continue
        parts.append(f"## {tab.get('name') or f'Sheet{i + 1}'}")
        for r in sorted(rows):
            line = " | ".join(t for _c, t in sorted(rows[r]))
            size += len(line) + 1
            if size > limit:
                parts.append("[truncated]")
                return "\n".join(parts)
            parts.append(line)
    return "\n".join(parts)


def preview_tabs(tabs: list[dict], max_rows: int = 200, max_cols: int = 50) -> tuple[list[dict], bool]:
    """The top-left corner of each tab, for a read-only preview."""
    out: list[dict] = []
    truncated = False
    for tab in tabs:
        cells = []
        for cell in tab.get("celldata") or []:
            if cell.get("r", 0) < max_rows and cell.get("c", 0) < max_cols:
                cells.append(cell)
            else:
                truncated = True
        config = dict(tab.get("config") or {})
        config.pop("borderInfo", None)
        out.append({k: v for k, v in tab.items() if k not in ("celldata", "config")} | {"celldata": cells, "config": config})
    return out, truncated


def cell_count(data: Any) -> int:
    return sum(len(_grid(t)) for t in _tabs(data))
