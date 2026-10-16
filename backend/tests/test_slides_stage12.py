"""Slides (Stage 12): "Generate a deck from these files" — slides drafted
from the files someone can read, each naming its sources; the model is
stubbed. And faithful Office previews: with the converter configured, a
PowerPoint file becomes a PDF made from its very bytes (the converter is
stubbed too), and a legacy .ppt is accepted on upload."""
from __future__ import annotations

import json

import pytest

from tests.test_drive_index_wp51 import _auth, _upload, _user, client  # noqa: F401 — fixtures
from tests.test_editor_ai_wp55 import _m, _text_of, model  # noqa: F401


@pytest.mark.asyncio
async def test_a_deck_is_drafted_from_files_with_its_sources(client, db, user, provider, model):  # noqa: F811
    deck = await _upload(client, user, "talk.txt", b"(the presentation)")
    budget = await _upload(client, user, "budget.txt", b"We spent $4,200 on the roof in March.")
    plan = await _upload(client, user, "plan.txt", b"Next year: solar panels, then the garden.")
    model["replies"]["You draft a slide deck"] = "Here you go:\n```json\n" + json.dumps({"slides": [
        {"title": "House update", "bullets": [], "notes": "Welcome.", "sources": []},
        {"title": "The roof", "bullets": ["$4,200 in March"], "notes": "The roof was the big one.", "sources": ["S1"]},
        {"title": "What's next", "bullets": ["Solar panels", "Then the garden"], "notes": "Two plans.", "sources": ["S2", "S9"]},
        {"title": "", "bullets": [], "notes": "empty — dropped"},
    ]}) + "\n```"
    r = await client.post(f"/api/files/{deck['id']}/ai/deck", headers=_auth(user), json={
        **_m(provider), "file_ids": [budget["id"], plan["id"]], "brief": "For the family meeting", "slides": 5,
    })
    assert r.status_code == 200, r.text
    got = r.json()
    assert [s["title"] for s in got["slides"]] == ["House update", "The roof", "What's next"]
    assert got["slides"][1]["sources"] == [{"file_id": budget["id"], "name": "budget.txt"}]
    assert [s["name"] for s in got["slides"][2]["sources"]] == ["plan.txt"]  # S9 isn't a source
    sent = _text_of(model["calls"][0]["messages"][0])
    assert '<source id="S1" name="budget.txt">' in sent and "solar panels" in sent and "family meeting" in sent

    # Someone else's file can't be a source.
    sam = await _user(db, "sam")
    theirs = await _upload(client, sam, "secret.txt", b"private")
    r = await client.post(f"/api/files/{deck['id']}/ai/deck", headers=_auth(user), json={
        **_m(provider), "file_ids": [theirs["id"]],
    })
    assert r.status_code == 404


@pytest.mark.asyncio
async def test_powerpoint_preview_is_converted_and_follows_new_bytes(client, db, user, monkeypatch, tmp_path):
    import io

    from pptx import Presentation

    from app.files import convert

    monkeypatch.setenv("CONVERT_URL", "http://convert.test:3000")
    sent: list[str] = []

    async def fake_post(path, filename):
        sent.append(filename)
        return b"%PDF-1.4 fake " + path.read_bytes()[:10]

    monkeypatch.setattr(convert, "_post", fake_post)

    async def enqueue(file_id):
        await convert.convert(file_id)

    monkeypatch.setattr("app.tasks.queue.enqueue_office_convert", enqueue)

    buf = io.BytesIO()
    p = Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Hello deck"
    p.save(buf)
    f = await _upload(client, user, "Holiday plan.pptx", buf.getvalue(), mime="application/octet-stream")

    r = await client.get(f"/api/files/{f['id']}/text", headers=_auth(user))
    assert r.json()["preview"] == "ready"  # asked → converted (inline here)
    assert sent == ["file.pptx"]  # never the person's file name
    pdf = await client.get(f"/api/files/{f['id']}/office-preview", headers=_auth(user))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"

    sam = await _user(db, "sam")
    assert (await client.get(f"/api/files/{f['id']}/office-preview", headers=_auth(sam))).status_code == 404

    # Off: no preview offered.
    monkeypatch.delenv("CONVERT_URL")
    assert (await client.get(f"/api/files/{f['id']}/text", headers=_auth(user))).json()["preview"] is None
    assert (await client.get(f"/api/files/{f['id']}/office-preview", headers=_auth(user))).status_code == 404


@pytest.mark.asyncio
async def test_a_legacy_ppt_is_accepted(client, db, user):
    # An OLE compound file with PowerPoint's record header where
    # ``filetype`` looks for it.
    ole = bytearray(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504)
    ole += b"\x0f\x00\xe8\x03" + b"\x00" * 2044
    f = await _upload(client, user, "old talk.ppt", bytes(ole), mime="application/vnd.ms-powerpoint")
    assert f["filename"] == "old talk.ppt"
