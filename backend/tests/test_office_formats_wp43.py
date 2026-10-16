"""Office files (WP-4.3): uploads behind a zip guard, text read on the
worker, "Open in Promptly Docs" (import + seed once), and exports rendered
from the live state with images inlined."""
from __future__ import annotations

import base64
import io
import uuid
import zipfile

import httpx
import pytest


@pytest.fixture
async def client(db, tmp_path, monkeypatch):
    from app import rate_limit
    from app.files import storage
    from app.main import app

    monkeypatch.setattr(storage, "UPLOAD_ROOT", tmp_path.resolve())
    # The worker isn't running in tests — read Office text in-process.
    from app.tasks import queue

    calls: list = []

    async def inline(file_id):
        from app.files.office_formats import extract_and_store

        calls.append(file_id)
        await extract_and_store(file_id)

    monkeypatch.setattr(queue, "enqueue_office_text", inline)
    rate_limit._settings.RATE_LIMIT_ENABLED = False
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as c:
            c.office_calls = calls  # type: ignore[attr-defined]
            yield c
    finally:
        rate_limit._settings.RATE_LIMIT_ENABLED = True


def _auth(user):
    from app.auth.utils import create_access_token

    return {"Authorization": f"Bearer {create_access_token(user.id, token_version=user.token_version)}"}


def _docx(*paragraphs: str, table=None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def _upload(client, user, name, data, mime="application/octet-stream"):
    return await client.post(
        "/api/files/",
        files={"file": (name, data, mime)},
        data={"scope": "mine"},
        headers=_auth(user),
    )


# ---------------------------------------------------------------- uploads


@pytest.mark.asyncio
async def test_a_word_file_uploads_and_its_text_is_searchable(client, db, user):
    from app.files.models import UserFile

    r = await _upload(client, user, "Lease.docx", _docx("Tenant agrees to quarterly inspections.",
                                                         table=[["Rent", "1450"]]))
    assert r.status_code == 201, r.text
    f = r.json()
    assert f["mime_type"].endswith("wordprocessingml.document")
    assert client.office_calls == [uuid.UUID(f["id"])]  # read off the request path
    row = await db.get(UserFile, uuid.UUID(f["id"]), populate_existing=True)
    assert "quarterly inspections" in row.content_text and "Rent | 1450" in row.content_text
    hits = (await client.get("/api/files/search?q=inspections", headers=_auth(user))).json()["hits"]
    assert any(h["file"]["id"] == f["id"] for h in hits)
    text = (await client.get(f"/api/files/{f['id']}/text", headers=_auth(user))).json()
    assert "quarterly" in text["text"] and text["pending"] is False


@pytest.mark.asyncio
async def test_excel_and_powerpoint_text(client, db, user, tmp_path):
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches

    wb = openpyxl.Workbook()
    wb.active.title = "Budget"
    wb.active.append(["Groceries", 600])
    xbuf = io.BytesIO()
    wb.save(xbuf)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly review"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Revenue up 12%"
    pbuf = io.BytesIO()
    prs.save(pbuf)

    from app.files.models import UserFile

    for name, data, expect in (("Budget.xlsx", xbuf.getvalue(), "Groceries | 600"),
                               ("Review.pptx", pbuf.getvalue(), "Revenue up 12%")):
        r = await _upload(client, user, name, data)
        assert r.status_code == 201, r.text
        row = await db.get(UserFile, uuid.UUID(r.json()["id"]), populate_existing=True)
        assert expect in (row.content_text or ""), (name, row.content_text)


@pytest.mark.asyncio
async def test_a_zip_bomb_and_a_fake_word_file_are_refused(client, db, user):
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "<w:document>" + "A" * (40 * 1024 * 1024) + "</w:document>")
    r = await _upload(client, user, "Huge.docx", bomb.getvalue())
    assert r.status_code == 400, r.text
    assert "expands" in r.text.lower()

    fake = io.BytesIO()
    with zipfile.ZipFile(fake, "w") as zf:
        zf.writestr("hello.txt", "not a document")
    r = await _upload(client, user, "Fake.docx", fake.getvalue())
    assert r.status_code == 400, r.text


# ---------------------------------------------------------------- open as document


def _yjs_paragraph(text: str) -> str:
    from pycrdt import Doc, XmlElement, XmlFragment, XmlText

    doc = Doc()
    frag = doc.get("default", type=XmlFragment)
    frag.children.append(XmlElement("paragraph", {}, [XmlText(text)]))
    return base64.b64encode(doc.get_update()).decode()


@pytest.mark.asyncio
async def test_open_in_promptly_docs_makes_a_linked_copy_seeded_once(client, db, user):
    from app.files.models import FileFolder, UserFile

    folder = FileFolder(user_id=user.id, name="Contracts")
    db.add(folder)
    await db.commit()
    r = await client.post("/api/files/", files={"file": ("Offer.docx", _docx("Welcome aboard."), "application/octet-stream")},
                          data={"scope": "mine", "folder_id": str(folder.id)}, headers=_auth(user))
    src = r.json()

    r = await client.post("/api/documents/import", json={"source_file_id": src["id"]}, headers=_auth(user))
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["filename"] == "Offer.html" and doc["folder_id"] == str(folder.id)
    row = await db.get(UserFile, uuid.UUID(doc["id"]), populate_existing=True)
    assert row.source_file_id == uuid.UUID(src["id"])

    r = await client.post(f"/api/documents/{doc['id']}/initial-state",
                          json={"yjs_state": _yjs_paragraph("Welcome aboard.")}, headers=_auth(user))
    assert r.status_code == 200, r.text
    row = await db.get(UserFile, uuid.UUID(doc["id"]), populate_existing=True)
    assert "Welcome aboard." in (row.content_text or "")
    # Never seeded twice.
    again = await client.post(f"/api/documents/{doc['id']}/initial-state",
                              json={"yjs_state": _yjs_paragraph("Again")}, headers=_auth(user))
    assert again.status_code == 409
    # Only a .docx can be opened this way.
    r = await _upload(client, user, "notes.txt", b"plain words", "text/plain")
    bad = await client.post("/api/documents/import", json={"source_file_id": r.json()["id"]}, headers=_auth(user))
    assert bad.status_code == 422


# ---------------------------------------------------------------- exports (R-219)


@pytest.mark.asyncio
async def test_exports_follow_the_live_state_and_carry_images(client, db, user):
    from app.files.models import DocumentState, UserFile
    from app.files.storage import absolute_path

    r = await client.post("/api/documents", json={"scope": "mine", "name": "Report"}, headers=_auth(user))
    doc = r.json()
    # An image asset of this document.
    png = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (4, 4), "red").save(png, format="PNG")
    r = await client.post(f"/api/documents/{doc['id']}/assets", files={"file": ("dot.png", png.getvalue(), "image/png")},
                          headers=_auth(user))
    assert r.status_code == 201, r.text
    img_url = r.json()["url"]

    from pycrdt import Doc, XmlElement, XmlFragment, XmlText

    ydoc = Doc()
    frag = ydoc.get("default", type=XmlFragment)
    frag.children.append(XmlElement("paragraph", {}, [XmlText("Live words")]))
    frag.children.append(XmlElement("image", {"src": img_url}))
    task_list = XmlElement("taskList", {}, [
        XmlElement("taskItem", {"checked": True}, [XmlElement("paragraph", {}, [XmlText("Done thing")])]),
        XmlElement("taskItem", {"checked": False}, [XmlElement("paragraph", {}, [XmlText("Open thing")])]),
    ])
    frag.children.append(task_list)
    ds = await db.get(DocumentState, uuid.UUID(doc["id"]))
    ds.yjs_update = ydoc.get_update()
    ds.version = 3
    await db.commit()
    # The on-disk snapshot is stale — the export must not use it.
    row = await db.get(UserFile, uuid.UUID(doc["id"]))
    absolute_path(row.storage_path).write_text("<p>Old snapshot</p>", encoding="utf-8")

    html = (await client.get(f"/api/documents/{doc['id']}/download?format=html", headers=_auth(user))).text
    assert html.startswith("<!DOCTYPE html>") and "<meta charset=\"utf-8\"" in html
    assert "Live words" in html and "Old snapshot" not in html
    assert "data:image/png;base64," in html and "/assets/" not in html

    md = (await client.get(f"/api/documents/{doc['id']}/download?format=md", headers=_auth(user))).text
    assert "[x] Done thing" in md and "[ ] Open thing" in md

    frag_html = (await client.get(f"/api/documents/{doc['id']}/download?format=fragment", headers=_auth(user))).text
    assert "<!DOCTYPE" not in frag_html and "Live words" in frag_html

    pdf = await client.get(f"/api/documents/{doc['id']}/download?format=pdf", headers=_auth(user))
    assert pdf.status_code == 200 and b"/Subtype /Image" in pdf.content


@pytest.mark.asyncio
async def test_a_word_attachment_reads_as_its_text_in_chat(client, db, user):
    from app.files.models import UserFile
    from app.files.prompt import build_attachment_preamble

    r = await _upload(client, user, "Brief.docx", _docx("The launch date is 3 March."))
    row = await db.get(UserFile, uuid.UUID(r.json()["id"]), populate_existing=True)
    assert "The launch date is 3 March." in build_attachment_preamble([row])
