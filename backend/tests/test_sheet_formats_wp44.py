"""Spreadsheets in and out (WP-4.4): .xlsx / .ods / .csv come in with their
formatting, "Open in Promptly Sheets" makes a linked copy, uploaded
spreadsheets preview as a grid, and downloads carry the look back out."""
from __future__ import annotations

import io
import uuid
import zipfile

import httpx
import pytest


@pytest.fixture
async def client(db, tmp_path, monkeypatch):
    from app import rate_limit
    from app.files import storage
    from app.main import app
    from app.tasks import queue

    monkeypatch.setattr(storage, "UPLOAD_ROOT", tmp_path.resolve())

    async def no_worker(file_id):  # Office text extraction isn't under test here
        return None

    monkeypatch.setattr(queue, "enqueue_office_text", no_worker)
    rate_limit._settings.RATE_LIMIT_ENABLED = False
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as c:
            yield c
    finally:
        rate_limit._settings.RATE_LIMIT_ENABLED = True


def _auth(user):
    from app.auth.utils import create_access_token

    return {"Authorization": f"Bearer {create_access_token(user.id, token_version=user.token_version)}"}


async def _other_user(db):
    from app.auth.models import User

    u = User(email=f"bob-{uuid.uuid4().hex[:6]}@example.com", username=f"bob{uuid.uuid4().hex[:6]}",
             password_hash="x", role="user")
    db.add(u)
    await db.commit()
    await db.refresh(u)
    return u


def _budget_xlsx() -> bytes:
    """A small but real-looking workbook: a merged title, a coloured header
    row with a border, money and dates, a total formula, a frozen header,
    a web link, a comment and a cross-tab formula."""
    import datetime

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.colors import Color

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Quarterly budget"
    ws["A1"].font = Font(bold=True, size=14, color="FF1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:C1")
    ws.append(["Item", "Cost", "Due"])
    for ref in ("A2", "B2", "C2"):
        # A theme colour (accent 4, lighter) — what Excel writes for most fills.
        ws[ref].fill = PatternFill("solid", fgColor=Color(theme=7, tint=0.6))
        ws[ref].font = Font(bold=True)
        ws[ref].border = Border(bottom=Side(style="medium", color="FF000000"))
    ws.append(["Rent", 1200.5, datetime.date(2026, 3, 31)])
    ws.append(["Power", 180, datetime.date(2026, 4, 15)])
    ws.append(["Total", "=SUM(B3:B4)", None])
    ws["B3"].number_format = "#,##0.00"
    ws["C3"].number_format = "yyyy-mm-dd"
    ws["C4"].number_format = "yyyy-mm-dd"
    ws["A5"].font = Font(italic=True)
    ws["A6"] = "Details"
    ws["A6"].hyperlink = "https://example.com/budget"
    ws["A4"].comment = Comment("Check the bill", "Sam")
    ws["A7"] = "00123"  # a code, not a number
    ws.column_dimensions["A"].width = 28
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A3"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Twice the total"
    notes["B1"] = "=Budget!B5*2"
    hidden = wb.create_sheet("Lookup")
    hidden["A1"] = "x"
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(tab: dict, r: int, c: int) -> dict:
    return next(x["v"] for x in tab["celldata"] if x["r"] == r and x["c"] == c)


def _ods() -> bytes:
    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
  office:version="1.3">
 <office:automatic-styles>
  <style:style style:name="co1" style:family="table-column"><style:table-column-properties style:column-width="2in"/></style:style>
  <style:style style:name="ce1" style:family="table-cell">
   <style:table-cell-properties fo:background-color="#ffe699"/>
   <style:text-properties fo:font-weight="bold"/>
  </style:style>
 </office:automatic-styles>
 <office:body><office:spreadsheet>
  <table:table table:name="Plan">
   <table:table-column table:style-name="co1"/>
   <table:table-column table:number-columns-repeated="1023"/>
   <table:table-row>
    <table:table-cell table:style-name="ce1" office:value-type="string" table:number-columns-spanned="2"><text:p>Header</text:p></table:table-cell>
    <table:covered-table-cell/>
   </table:table-row>
   <table:table-row>
    <table:table-cell office:value-type="float" office:value="2"><text:p>2</text:p></table:table-cell>
    <table:table-cell office:value-type="percentage" office:value="0.25"><text:p>25%</text:p></table:table-cell>
   </table:table-row>
   <table:table-row>
    <table:table-cell office:value-type="float" office:value="3"><text:p>3</text:p></table:table-cell>
    <table:table-cell office:value-type="date" office:date-value="2026-03-31"><text:p>31/03/26</text:p></table:table-cell>
   </table:table-row>
   <table:table-row>
    <table:table-cell table:formula="of:=SUM([.A2:.A3])" office:value-type="float" office:value="5"><text:p>5</text:p></table:table-cell>
    <table:table-cell table:formula="of:=IF([$Data.B2]&gt;1;&quot;big;ish&quot;;&quot;small&quot;)" office:value-type="string"><text:p>big;ish</text:p></table:table-cell>
   </table:table-row>
   <table:table-row table:number-rows-repeated="1048572"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
  </table:table>
  <table:table table:name="Data"><table:table-row><table:table-cell/><table:table-cell office:value-type="float" office:value="7"><text:p>7</text:p></table:table-cell></table:table-row></table:table>
 </office:spreadsheet></office:body>
</office:document-content>"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", content)
    return buf.getvalue()


# ---------------------------------------------------------------- display text


def test_number_formats_show_what_excel_shows():
    from app.workspaces.sheet_numfmt import format_value as f

    assert f(1200.5, "#,##0.00") == "1,200.50"
    assert f(46112, "yyyy-mm-dd") == "2026-03-31"
    assert f(46112, "d-mmm-yy") == "31-Mar-26"
    assert f(0.256, "0.0%") == "25.6%"
    assert f(-1234.5, "#,##0.00;(#,##0.00)") == "(1,234.50)"
    assert f(1234.5, "[$€-407]#,##0.00") == "€1,234.50"
    assert f(0, '#,##0;-#,##0;"-"') == "-"
    assert f(0.5, "h:mm AM/PM") == "12:00 PM"
    assert f(0.1 + 0.2, "General") == "0.3"
    # Beyond the renderer → None (the editor formats it with SSF).
    assert f(5.25, "# ?/?") is None
    assert f(0.75, "[h]:mm:ss") is None


# ---------------------------------------------------------------- xlsx in


def test_an_xlsx_comes_in_looking_like_itself():
    from app.workspaces.sheet_io import import_spreadsheet

    tabs, notes = import_spreadsheet(_budget_xlsx(), "Budget.xlsx")
    budget, other, lookup = tabs
    cfg = budget["config"]
    assert cfg["merge"] == {"0_0": {"r": 0, "c": 0, "rs": 1, "cs": 3}}
    assert cfg["columnlen"]["0"] == 201 and cfg["rowlen"]["0"] == 40
    assert budget["frozen"] == {"type": "rangeRow", "range": {"row_focus": 1, "column_focus": 0}}
    assert {"rangeType": "cell", "value": {"row_index": 1, "col_index": 0, "b": {"style": 8, "color": "#000000"}}} in cfg["borderInfo"]

    title = _cell(budget, 0, 0)
    assert title["bl"] == 1 and title["fs"] == 14 and title["fc"] == "#1F4E79" and title["ht"] == 0
    assert title["mc"] == {"r": 0, "c": 0, "rs": 1, "cs": 3}
    assert _cell(budget, 0, 1)["mc"] == {"r": 0, "c": 0}
    header = _cell(budget, 1, 0)
    assert header["bl"] == 1 and header["bg"].startswith("#") and header["bg"] != "#FFC000"  # tinted accent 4

    rent = _cell(budget, 2, 1)
    assert rent["v"] == 1200.5 and rent["m"] == "1,200.50" and rent["ct"] == {"fa": "#,##0.00", "t": "n"}
    assert rent["ht"] == 2  # numbers right-aligned, as Excel shows them
    assert "ht" not in _cell(budget, 2, 0)
    due = _cell(budget, 2, 2)
    assert due["v"] == 46112 and due["m"] == "2026-03-31" and due["ct"]["t"] == "d"
    total = _cell(budget, 4, 1)
    assert total["f"] == "=SUM(B3:B4)" and "m" not in total  # no cached result: the editor works it out
    assert _cell(budget, 4, 0)["it"] == 1
    link = _cell(budget, 5, 0)
    assert link["hl"]["r"] == 5 and budget["hyperlink"]["5_0"] == {"linkType": "webpage", "linkAddress": "https://example.com/budget"}
    assert _cell(budget, 6, 0)["v"] == "00123"

    assert _cell(other, 0, 1)["f"] == "=Budget!B5*2"
    assert lookup.get("hide") == 1
    assert notes == ["1 cell comment wasn't brought over."]


def test_a_formula_without_a_result_keeps_its_number_format():
    from app.workspaces.sheet_io import _fortune_cell

    cell = _fortune_cell(None, "=SUM(B3:B4)", "#,##0.00")
    assert cell == {"v": None, "ct": {"fa": "#,##0.00", "t": "n"}, "f": "=SUM(B3:B4)"}


def test_cached_formula_results_come_in_with_the_formula():
    """A workbook saved by Excel carries each formula's last result."""
    import re as _re

    from app.workspaces.sheet_io import import_spreadsheet

    raw = _budget_xlsx()
    src = zipfile.ZipFile(io.BytesIO(raw))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as out:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                text = _re.sub(r"(<c r=\"B5\"[^>]*>)<f>SUM\(B3:B4\)</f>(<v>\s*</v>|<v/>)?", r"\1<f>SUM(B3:B4)</f><v>1380.5</v>", text)
                assert "<v>1380.5</v>" in text
                data = text.encode()
            out.writestr(item, data)
    tabs, _notes = import_spreadsheet(buf.getvalue(), "Budget.xlsx")
    total = _cell(tabs[0], 4, 1)
    assert total["f"] == "=SUM(B3:B4)" and total["v"] == 1380.5 and total["m"] == "1380.5"


def test_a_zip_bomb_or_a_fake_workbook_is_refused():
    from app.workspaces.sheet_io import SheetImportError, import_spreadsheet

    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", b"<workbook/>")
        zf.writestr("xl/big.bin", b"\0" * (20 * 1024 * 1024))
    with pytest.raises(SheetImportError):
        import_spreadsheet(bomb.getvalue(), "big.xlsx")
    with pytest.raises(SheetImportError):
        import_spreadsheet(b"not a zip at all", "fake.xlsx")


# ---------------------------------------------------------------- ods / csv in


def test_an_ods_comes_in_with_formulas_spans_and_styles():
    from app.workspaces.sheet_io import import_spreadsheet

    tabs, _notes = import_spreadsheet(_ods(), "Plan.ods")
    plan, data = tabs
    assert plan["name"] == "Plan" and data["name"] == "Data"
    head = _cell(plan, 0, 0)
    assert head["v"] == "Header" and head["bl"] == 1 and head["bg"] == "#FFE699"
    assert plan["config"]["merge"]["0_0"] == {"r": 0, "c": 0, "rs": 1, "cs": 2}
    assert plan["config"]["columnlen"]["0"] == 192
    assert _cell(plan, 1, 1)["ct"]["fa"] == "0.00%" and _cell(plan, 1, 1)["m"] == "25%"
    due = _cell(plan, 2, 1)
    assert due["v"] == 46112 and due["ct"]["t"] == "d"
    total = _cell(plan, 3, 0)
    assert total["f"] == "=SUM(A2:A3)" and total["v"] == 5
    # Argument separators become commas — but not inside a string.
    assert _cell(plan, 3, 1)["f"] == '=IF(Data!B2>1,"big;ish","small")'
    # A million repeated empty rows don't become a million rows.
    assert max(c["r"] for c in plan["celldata"]) == 3
    assert _cell(data, 0, 1)["v"] == 7


def test_csv_keeps_codes_with_leading_zeros_as_text():
    from app.workspaces.sheet_io import from_csv

    tab = from_csv(b"code,amount\n00123,12.5\n", "Codes")[0]
    assert _cell(tab, 1, 0)["v"] == "00123" and _cell(tab, 1, 1)["v"] == 12.5


# ---------------------------------------------------------------- xlsx out


def test_a_sheet_downloads_with_its_formatting():
    from openpyxl import load_workbook

    from app.workspaces.sheet_io import to_xlsx

    data = [
        {"name": "Main", "order": 0, "frozen": {"type": "rangeBoth", "range": {"row_focus": 0, "column_focus": 0}},
         "hyperlink": {"2_0": {"linkType": "webpage", "linkAddress": "https://example.com"}},
         "config": {
             "merge": {"0_0": {"r": 0, "c": 0, "rs": 1, "cs": 2}},
             "columnlen": {"0": 152}, "rowlen": {"0": 40}, "colhidden": {"3": 0},
             "borderInfo": [{"rangeType": "range", "borderType": "border-outside", "style": "1", "color": "#ff0000",
                             "range": [{"row": [1, 2], "column": [0, 1]}]}],
         },
         "celldata": [
             {"r": 0, "c": 0, "v": {"v": "Title", "m": "Title", "bl": 1, "fs": 16, "fc": "rgb(31, 78, 121)", "ht": 0, "bg": "#fff2cc"}},
             {"r": 1, "c": 0, "v": {"v": 46112, "m": "2026-03-31", "ct": {"fa": "yyyy-mm-dd", "t": "d"}}},
             {"r": 1, "c": 1, "v": {"v": "1200.5", "m": "1,200.50", "ct": {"fa": "#,##0.00", "t": "n"}, "it": 1, "un": 1}},
             {"r": 2, "c": 0, "v": {"v": "Site", "m": "Site", "hl": {"r": 2, "c": 0, "id": "x"}}},
             {"r": 2, "c": 1, "v": {"v": 1200.5, "m": "1200.5", "f": "=B2"}},
             {"r": 3, "c": 0, "v": {"m": "rich", "ct": {"t": "inlineStr", "s": [{"v": "ri"}, {"v": "ch", "bl": 1}]}}},
         ]},
        {"name": "Old", "order": 1, "hide": 1, "celldata": [{"r": 0, "c": 0, "v": {"v": "x"}}]},
    ]
    wb = load_workbook(io.BytesIO(to_xlsx(data)))
    ws = wb["Main"]
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:B1"]
    assert ws.column_dimensions["A"].width == 21 and ws.row_dimensions[1].height == 30
    assert ws.column_dimensions["D"].hidden
    assert ws.freeze_panes == "B2"
    assert ws["A1"].font.b and ws["A1"].font.sz == 16 and ws["A1"].font.color.rgb == "FF1F4E79"
    assert ws["A1"].fill.fgColor.rgb == "FFFFF2CC" and ws["A1"].alignment.horizontal == "center"
    assert ws["A2"].number_format == "yyyy-mm-dd" and ws["A2"].value.year == 2026
    assert ws["B2"].value == 1200.5 and ws["B2"].number_format == "#,##0.00" and ws["B2"].font.i and ws["B2"].font.u == "single"
    assert ws["A2"].border.left.style == "thin" and ws["A2"].border.top.style == "thin"
    assert ws["A2"].border.right is None or ws["A2"].border.right.style is None
    assert ws["B3"].border.right.style == "thin" and ws["B3"].border.bottom.style == "thin"
    assert ws["A3"].hyperlink.target == "https://example.com"
    assert ws["B3"].value == "=B2"
    assert ws["A4"].value == "rich"
    assert wb["Old"].sheet_state == "hidden"
    assert wb.calculation.fullCalcOnLoad


def test_an_imported_workbook_round_trips():
    from openpyxl import load_workbook

    from app.workspaces.sheet_io import import_spreadsheet, to_xlsx

    tabs, _ = import_spreadsheet(_budget_xlsx(), "Budget.xlsx")
    ws = load_workbook(io.BytesIO(to_xlsx(tabs)))["Budget"]
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:C1"]
    assert ws.column_dimensions["A"].width == 28 and ws.row_dimensions[1].height == 30
    assert ws.freeze_panes == "A3"
    assert ws["A1"].font.b and ws["A1"].font.color.rgb == "FF1F4E79"
    assert ws["A2"].border.bottom.style == "medium"
    assert ws["B3"].number_format == "#,##0.00" and ws["C3"].value.month == 3
    assert ws["B5"].value == "=SUM(B3:B4)" and ws["A7"].value == "00123"
    assert ws["A6"].hyperlink.target == "https://example.com/budget"


# ---------------------------------------------------------------- endpoints


@pytest.mark.asyncio
async def test_open_in_promptly_sheets_makes_a_linked_formatted_copy(client, db, user):
    from app.files.models import FileFolder, UserFile

    folder = FileFolder(user_id=user.id, name="Money")
    db.add(folder)
    await db.commit()
    r = await client.post("/api/files/", files={"file": ("Budget.xlsx", _budget_xlsx(), "application/octet-stream")},
                          data={"scope": "mine", "folder_id": str(folder.id)}, headers=_auth(user))
    assert r.status_code in (200, 201), r.text
    src = r.json()

    r = await client.post("/api/files/office/import-sheet", json={"source_file_id": src["id"]}, headers=_auth(user))
    assert r.status_code == 201, r.text
    body = r.json()
    new = body["file"]
    assert body["path"] == f"/files/sheet/{new['id']}"
    assert body["notes"] == ["1 cell comment wasn't brought over."]
    assert new["folder_id"] == str(folder.id)
    row = await db.get(UserFile, uuid.UUID(new["id"]), populate_existing=True)
    assert row.source_file_id == uuid.UUID(src["id"])
    assert "Quarterly budget" in (row.content_text or "") and "1,200.50" in (row.content_text or "")

    sheet = (await client.get(f"/api/files/{new['id']}/sheet", headers=_auth(user))).json()
    assert sheet["data"][0]["config"]["merge"]["0_0"]["cs"] == 3
    # Download gives the look back.
    from openpyxl import load_workbook

    xlsx = await client.get(f"/api/files/{new['id']}/sheet/export", params={"format": "xlsx"}, headers=_auth(user))
    assert xlsx.status_code == 200
    assert load_workbook(io.BytesIO(xlsx.content))["Budget"]["A1"].font.b

    # The original is untouched, and someone without access gets nothing.
    orig = await db.get(UserFile, uuid.UUID(src["id"]), populate_existing=True)
    assert orig.filename == "Budget.xlsx" and orig.trashed_at is None
    bob = await _other_user(db)
    r = await client.post("/api/files/office/import-sheet", json={"source_file_id": src["id"]}, headers=_auth(bob))
    assert r.status_code == 404


@pytest.mark.asyncio
async def test_only_spreadsheets_open_as_sheets(client, db, user):
    r = await client.post("/api/files/", files={"file": ("notes.txt", b"hello", "text/plain")},
                          data={"scope": "mine"}, headers=_auth(user))
    r = await client.post("/api/files/office/import-sheet", json={"source_file_id": r.json()["id"]}, headers=_auth(user))
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_an_uploaded_spreadsheet_previews_as_a_grid(client, db, user):
    r = await client.post("/api/files/", files={"file": ("Plan.ods", _ods(), "application/octet-stream")},
                          data={"scope": "mine"}, headers=_auth(user))
    assert r.status_code in (200, 201), r.text
    fid = r.json()["id"]
    r = await client.get(f"/api/files/{fid}/sheet-preview", headers=_auth(user))
    assert r.status_code == 200, r.text
    body = r.json()
    assert [t["name"] for t in body["tabs"]] == ["Plan", "Data"]
    assert body["truncated"] is False
    bob = await _other_user(db)
    assert (await client.get(f"/api/files/{fid}/sheet-preview", headers=_auth(bob))).status_code == 404


@pytest.mark.asyncio
async def test_importing_into_a_sheet_takes_ods_and_reports_notes(client, db, user):
    r = await client.post("/api/files/office", json={"kind": "sheet", "title": "Scratch"}, headers=_auth(user))
    assert r.status_code == 201, r.text
    sid = r.json()["id"]
    r = await client.post(f"/api/files/{sid}/sheet/import",
                          files={"file": ("Plan.ods", _ods(), "application/octet-stream")}, headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["data"][0]["name"] == "Plan" and r.json()["notes"] == []
    r = await client.post(f"/api/files/{sid}/sheet/import",
                          files={"file": ("Budget.xlsx", _budget_xlsx(), "application/octet-stream")}, headers=_auth(user))
    assert r.json()["notes"] == ["1 cell comment wasn't brought over."]
